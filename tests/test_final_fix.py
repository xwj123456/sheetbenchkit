"""Actual-file regressions for the three final-review findings."""

import json
import os
import re
import subprocess
import sys
from dataclasses import replace

import pytest
from openpyxl import Workbook
from test_artifacts import candidate
from test_loaders import case, workbook_bytes
from test_variants import make_base

from sheetbenchkit import artifacts as a
from sheetbenchkit import models as m
from sheetbenchkit.contracts import decode_document, to_document, validate_case, validate_schema
from sheetbenchkit.examples.independent_runner import solve
from sheetbenchkit.grader import grade_suite
from sheetbenchkit.loaders import load_inputs
from sheetbenchkit.reference import evaluate_reference
from sheetbenchkit.report import exit_code
from sheetbenchkit.runner import make_task_envelope, run_suite
from sheetbenchkit.variants import generate_suite


def integer_json(document, literal):
    raw = json.dumps(document)
    return re.sub(r"([0-9]+)\.0\b", r"\1e0", raw) if literal == "1e0" else raw


def actual_candidate(root, *, width=1, count=2, format="csv", sparse=False, op="count"):
    root.mkdir()
    headers = ["amount"] + [f"c{i}" for i in range(1, width)]
    rows = (4, 5) if sparse else (2, count + 1)
    if format == "csv":
        raw = (",".join(headers) + "\n" + (",".join(["1"] * width) + "\n") * count).encode()
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Main"
        sheet.append(headers)
        for _ in range(1 if sparse else count):
            sheet.append([1] * width)
        book.save(root / "input.xlsx")
        book.close()
        raw = (root / "input.xlsx").read_bytes()
    metric = {"op": "count", "source": "main"} if op == "count" else None
    spec, manifest = case(root, raw, format=format, rows=rows, metric=metric)
    task = {"description": "Count declared rows or sum amount; obey table bounds."}
    confirmation = {
        "reviewed_contract_sha256": a.sha256_bytes(a.canonical_json(spec)),
        "task_sha256": a.sha256_bytes(a.canonical_json(task)),
    }
    manifest = replace(manifest, business_confirmation=confirmation)
    (root / "task.json").write_bytes(a.canonical_json(task))
    validated = validate_case(spec, manifest, root)
    if width > 100 or count > 1000:
        expected = m.ExpectedResult("ABSTAIN", None, None, "UNSUPPORTED_INPUT_LIMIT", ())
    elif sparse and op == "sum":
        expected = m.ExpectedResult("ABSTAIN", None, None, "MISSING_VALUE", ())
    else:
        value = str(2 if sparse else count) if op == "count" else f"{count}.00"
        binding = m.Binding(
            "main",
            "work",
            "@csv" if format == "csv" else "Main",
            1,
            rows,
            () if op == "count" else ("amount",),
        )
        expected = m.ExpectedResult(
            "VALUE", value, "count" if op == "count" else "CNY", None, (binding,)
        )
    return m.CandidateCase(
        validated,
        expected,
        "preset_independent",
        confirmation,
        {"method": "literal independent row/width/blank expectation"},
    )


def cloned(item, count):
    return tuple(
        replace(
            item,
            validated=replace(
                item.validated, manifest=replace(item.validated.manifest, case_id=f"C{i}")
            ),
        )
        for i in range(count)
    )


@pytest.mark.parametrize("count", [9, 10])
def test_rejected_csv_suite_budget_all_consumers(tmp_path, count):
    item = actual_candidate(tmp_path / "input", width=101, count=1000)
    assert not item.validated.inputs.tables
    assert a._read_cells(item.validated.inputs) == 101101
    items = cloned(item, count)
    root = tmp_path / "frozen"
    if count == 10:
        with pytest.raises(m.CaseError, match="SUITE_LIMIT"):
            a.freeze_suite(items, (), root)
        assert not root.exists()
    # Supplied caller bundles are not a trusted replacement for independent preflight.
    empty = replace(item.validated.inputs, tables={}, selected_read_cells={})
    untrusted = tuple(replace(c, validated=replace(c.validated, inputs=empty)) for c in items)
    suite = a.freeze_suite(untrusted, (), root)
    if count == 9:
        cases = a.preflight_suite(suite, root)
        assert sum(a._read_cells(c.validated.inputs) for c in cases) == 909909
        observations = tuple(
            m.Observation(
                c.gold.case_id,
                1,
                "x",
                "SUCCESS",
                json.dumps(
                    {
                        "schema_version": "1",
                        "case_id": c.gold.case_id,
                        "metric_id": "amount",
                        **to_document(item.expected),
                    }
                ),
                "",
                None,
                {},
            )
            for c in cases
        )
        report = grade_suite(
            suite, cases, m.ObservationBatch("1", (m.RunSpec("x", 1),), observations)
        )
        assert report.valid and report.status_counts["PASS"] == 9 and exit_code(report) == 0
        return
    with pytest.raises(m.CaseError, match="SUITE_LIMIT") as caught:
        a.preflight_suite(suite, root)
    assert caught.value.details["read_cells"] == 1011010
    marker = tmp_path / "starts"
    if os.name == "posix":
        with pytest.raises(m.CaseError, match="SUITE_LIMIT"):
            run_suite(
                suite,
                root,
                (
                    sys.executable,
                    "-c",
                    f"from pathlib import Path;Path({str(marker)!r}).write_text('started')",
                ),
            )
    assert not marker.exists()
    # Individually frozen cases preserve their exact hashes, unlike forged bundles.
    cases = []
    for index, c in enumerate(items):
        one_root = tmp_path / f"single{index}"
        one_suite = a.freeze_suite((c,), (), one_root)
        cases.extend(a.preflight_suite(one_suite, one_root))
    report = grade_suite(suite, tuple(cases), m.ObservationBatch("1", (m.RunSpec("x", 1),), ()))
    assert not report.valid and exit_code(report) == 2
    assert any(c.reason == "SUITE_LIMIT" for c in report.case_grades)


def test_omitted_xlsx_header_reads_are_retained(tmp_path):
    raw = workbook_bytes(headers=(1, "amount"), rows=(("a", 1), ("b", 2)))
    spec, manifest = case(tmp_path, raw, format="xlsx")
    bundle = load_inputs(spec, manifest, tmp_path)
    assert not bundle.tables
    assert a._read_cells(bundle) >= 2


@pytest.mark.parametrize("format", ["csv", "xlsx"])
@pytest.mark.parametrize("width,count", [(100, 2), (101, 2), (1, 1000), (1, 1001)])
def test_independent_table_boundaries(tmp_path, format, width, count):
    item = actual_candidate(tmp_path / "input", width=width, count=count, format=format)
    root = tmp_path / "frozen"
    suite = a.freeze_suite((item,), (), root)
    frozen = a.preflight_suite(suite, root)[0]
    envelope = to_document(make_task_envelope(frozen))
    envelope["inputs"][0]["relative_path"] = str(
        (root / "cases/loader" / frozen.validated.manifest.inputs[0].relative_path).resolve()
    )
    result = solve(envelope)
    assert result["status"] == item.expected.status
    assert result["value"] == item.expected.value
    assert result["reason"] == item.expected.reason


@pytest.mark.parametrize("op", ["count", "sum"])
def test_independent_sparse_trailing_xlsx_rows(tmp_path, op):
    item = actual_candidate(tmp_path / "input", format="xlsx", sparse=True, op=op)
    root = tmp_path / "frozen"
    suite = a.freeze_suite((item,), (), root)
    frozen = a.preflight_suite(suite, root)[0]
    envelope = to_document(make_task_envelope(frozen))
    envelope["inputs"][0]["relative_path"] = str(
        (root / "cases/loader" / frozen.validated.manifest.inputs[0].relative_path).resolve()
    )
    result = solve(envelope)
    assert (result["status"], result["value"], result["reason"]) == (
        item.expected.status,
        item.expected.value,
        item.expected.reason,
    )


@pytest.mark.parametrize("literal", ["1.0", "1e0"])
def test_integer_lexemes_load_reference_variant_freeze_grade(tmp_path, literal):
    base = make_base(tmp_path / "base", places=2)
    spec_doc = to_document(base.spec)
    spec_doc["sources"]["chosen"].update(header_row=1.0, rows=[2.0, 3.0])
    spec_doc["metric"]["places"] = 2.0
    spec = decode_document("contract", integer_json(spec_doc, literal))
    assert type(spec.sources["chosen"].header_row) is int
    assert all(type(v) is int for v in spec.sources["chosen"].rows)
    assert type(spec.metric.places) is int
    manifest_doc = to_document(base.manifest)
    for recipe in manifest_doc["variant_recipes"]:
        recipe.update(header_row=1.0, rows=[2.0, 3.0], row=float(recipe["row"]))
    manifest = decode_document("manifest", integer_json(manifest_doc, literal))
    assert all(type(r.row) is int and type(r.header_row) is int for r in manifest.variant_recipes)
    loaded = validate_case(spec, manifest, tmp_path / "base")
    assert evaluate_reference(loaded).value == "100.00"
    suite = generate_suite((loaded,), 7, tmp_path / "variants")
    cases = a.preflight_suite(suite, tmp_path / "variants")
    gold_doc = to_document(cases[0].gold)
    binding = gold_doc["expected"]["bindings"][0]
    binding.update(header_row=1.0, rows=[2.0, 3.0])
    gold = decode_document("gold", integer_json(gold_doc, literal))
    assert type(gold.expected.bindings[0].header_row) is int
    result_doc = {
        "schema_version": "1",
        "case_id": cases[0].gold.case_id,
        "metric_id": cases[0].validated.spec.metric_id,
        **gold_doc["expected"],
    }
    result = decode_document("result", integer_json(result_doc, literal))
    assert type(result.bindings[0].rows[0]) is int
    raw = (
        '{"schema_version":"1","runs":[{"config_id":"x","repeats":'
        + literal
        + '}],"observations":[{"case_id":"'
        + cases[0].gold.case_id
        + '","attempt":'
        + literal
        + ',"config_id":"x","execution_status":"SUCCESS",'
        '"raw_stdout":'
        + json.dumps(json.dumps(result_doc))
        + ',"raw_stderr":"","usage":{"cost":1.0},"runtime_metadata":{"elapsed":1.0}}]}'
    )
    batch = decode_document("observations", raw)
    assert type(batch.runs[0].repeats) is int and type(batch.observations[0].attempt) is int
    assert type(batch.observations[0].usage["cost"]) is float
    assert type(batch.observations[0].runtime_metadata["elapsed"]) is float
    report = grade_suite(suite, cases, batch)
    assert report.valid and report.status_counts["PASS"] == 1


@pytest.mark.parametrize("literal", ["1.0", "1e0"])
def test_actual_cli_integral_repeats_missing_observations(tmp_path, literal):
    root = tmp_path / "suite"
    a.freeze_suite((candidate(),), (), root)
    path = tmp_path / "observations.json"
    path.write_text(
        '{"schema_version":"1","runs":[{"config_id":"x","repeats":'
        + literal
        + '}],"observations":[]}'
    )
    output = tmp_path / "report"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "sheetbenchkit",
            "grade",
            "--suite",
            str(root),
            "--observations",
            str(path),
            "--output",
            str(output),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 1 and "Traceback" not in result.stderr
    report = json.loads((output / "report.json").read_text())
    assert report["valid"] and report["status_counts"]["NO_RESULT"] == 1


def test_direct_dataclass_integer_boundaries():
    objects = [
        m.SourceSpec("w", "@csv", 1.0, (2.0, 3.0)),
        m.Binding("s", "w", "@csv", 1.0, (2.0, 3.0), ()),
        m.MutationRecipe("target_numeric", "w", "@csv", 1.0, (2.0, 3.0), "v", 2.0, "1", None),
        m.RootSumSpec("sum", "s", "v", "CNY", places=2.0, rounding="HALF_UP"),
        m.RatioSpec("ratio", m.CountSpec("count", "s"), m.CountSpec("count", "s"), 2.0, "HALF_UP"),
        m.RunSpec("x", 1.0),
        m.Observation("a", 1.0, "x", "SUCCESS", "", "", None, {"elapsed": 1.0}),
    ]
    for obj in objects:
        for field in ("header_row", "places", "row", "attempt", "repeats"):
            if hasattr(obj, field):
                assert type(getattr(obj, field)) is int
        if hasattr(obj, "rows"):
            assert all(type(v) is int for v in obj.rows)


@pytest.mark.parametrize("bad", [True, 1.5, "1", None])
@pytest.mark.parametrize("slot", ["header_row", "rows", "places", "repeats", "attempt", "row"])
def test_invalid_integer_slots_controlled_refusal(tmp_path, bad, slot):
    base = make_base(tmp_path / "base")
    if slot in {"header_row", "rows", "places"}:
        name, doc = "contract", to_document(base.spec)
        if slot == "places":
            doc["metric"][slot] = bad
        else:
            doc["sources"]["chosen"][slot] = [bad, 3] if slot == "rows" else bad
    elif slot == "row":
        name, doc = "manifest", to_document(base.manifest)
        doc["variant_recipes"][0]["row"] = bad
    else:
        name = "observations"
        doc = {
            "schema_version": "1",
            "runs": [{"config_id": "x", "repeats": 1}],
            "observations": [
                {
                    "case_id": "x",
                    "attempt": 1,
                    "config_id": "x",
                    "execution_status": "SUCCESS",
                    "raw_stdout": "",
                    "raw_stderr": "",
                    "usage": None,
                    "runtime_metadata": {},
                }
            ],
        }
        doc["runs" if slot == "repeats" else "observations"][0][slot] = bad
    with pytest.raises(m.CaseError):
        decode_document(name, json.dumps(doc))
    # Public model ingress is also validated before downstream arithmetic/index consumers.
    if slot == "repeats":
        with pytest.raises(m.CaseError):
            validate_schema("observations", m.ObservationBatch("1", (m.RunSpec("x", bad),), ()))


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_row_limit_omission_accounts_only_selected_header(tmp_path, format):
    item = actual_candidate(tmp_path / "input", count=1001, width=2, format=format)
    assert not item.validated.inputs.tables
    assert a._read_cells(item.validated.inputs) == 2


def test_xlsx_missing_header_omission_retains_selected_data_reads(tmp_path):
    raw = workbook_bytes(headers=(), rows=(("a", 1), ("b", 2)))
    spec, manifest = case(tmp_path, raw, format="xlsx", boundaries=("MISSING_FIELD",))
    bundle = load_inputs(spec, manifest, tmp_path)
    assert not bundle.tables
    assert a._read_cells(bundle) == 4


def test_exact_million_selected_cells_is_legal(tmp_path):
    item = actual_candidate(tmp_path / "input", width=100, count=999)
    assert a._read_cells(item.validated.inputs) == 100000
    root = tmp_path / "frozen"
    suite = a.freeze_suite(cloned(item, 10), (), root)
    cases = a.preflight_suite(suite, root)
    assert sum(a._read_cells(c.validated.inputs) for c in cases) == 1000000
    report = grade_suite(suite, cases, m.ObservationBatch("1", (m.RunSpec("x", 1),), ()))
    assert report.valid and report.missing_count == 10


@pytest.mark.parametrize("name", ["gold", "result"])
@pytest.mark.parametrize("slot", ["header_row", "rows"])
@pytest.mark.parametrize("bad", [True, 1.5, "1"])
def test_binding_integer_wrong_types_are_controlled(name, slot, bad, tmp_path):
    root = tmp_path / "suite"
    suite = a.freeze_suite((candidate(),), (), root)
    frozen = a.preflight_suite(suite, root)[0]
    doc = to_document(frozen.gold)
    if name == "result":
        doc = {
            "schema_version": "1",
            "case_id": frozen.gold.case_id,
            "metric_id": frozen.validated.spec.metric_id,
            **doc["expected"],
        }
    binding = (doc["expected"] if name == "gold" else doc)["bindings"][0]
    binding[slot] = [bad, 3] if slot == "rows" else bad
    with pytest.raises(m.CaseError):
        decode_document(name, json.dumps(doc))


@pytest.mark.parametrize("literal", ["1.0000000000000001", "1.00000000000000000000001"])
def test_nonintegral_integer_lexemes_never_round_into_valid_repeats(literal):
    raw = (
        '{"schema_version":"1","runs":[{"config_id":"x","repeats":'
        + literal
        + '}],"observations":[]}'
    )
    with pytest.raises(m.CaseError, match="INVALID_OBSERVATIONS"):
        decode_document("observations", raw)
