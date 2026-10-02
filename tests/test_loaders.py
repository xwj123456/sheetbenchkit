"""Real CSV/OOXML fixtures catch type loss, clipping, malformed input and path escapes."""

import dataclasses
import hashlib
import importlib
import importlib.util
import io
import json
import re
import zipfile
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from sheetbenchkit import contracts
from sheetbenchkit.models import CaseError


def load(spec, manifest, root):
    assert importlib.util.find_spec("sheetbenchkit.loaders"), "Missing T2 bounded loaders"
    return importlib.import_module("sheetbenchkit.loaders").load_inputs(spec, manifest, root)


def case(tmp_path, data, *, format="csv", rows=(2, 3), sheet=None, boundaries=(), metric=None):
    path = tmp_path / f"input.{format}"
    path.write_bytes(data)
    spec_doc = {
        "schema_version": "1",
        "metric_id": "amount",
        "sources": {
            "main": {
                "file_id": "work",
                "sheet": sheet or ("@csv" if format == "csv" else "Main"),
                "header_row": 1,
                "rows": rows,
            }
        },
        "metric": metric
        or {
            "op": "sum",
            "source": "main",
            "field": "amount",
            "unit": "CNY",
            "places": 2,
            "rounding": "HALF_UP",
        },
    }
    manifest_doc = {
        "schema_version": "1",
        "case_id": "loader",
        "category": "custom",
        "inputs": [
            {
                "file_id": "work",
                "relative_path": path.name,
                "format": format,
                "sha256": hashlib.sha256(data).hexdigest(),
                "expected_missing": False,
            }
        ],
        "unit_labels": {"work": {spec_doc["sources"]["main"]["sheet"]: {"amount": "CNY"}}},
        "intentional_boundaries": boundaries,
        "business_confirmation": {},
        "variant_recipes": [],
    }
    return (
        contracts.decode_document("contract", json.dumps(spec_doc)),
        contracts.decode_document("manifest", json.dumps(manifest_doc)),
    )


def workbook_bytes(headers=("name", "amount"), rows=(("a", 0), ("b", 2))):
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


def rewrite_zip(data, replacements=None, extra=()):
    stream = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as target,
    ):
        for info in source.infolist():
            raw = source.read(info)
            if replacements and info.filename in replacements:
                raw = replacements[info.filename](raw)
            target.writestr(info.filename, raw)
        for name, raw in extra:
            target.writestr(name, raw)
    return stream.getvalue()


def reasons(bundle):
    return [item["reason"] for item in bundle.business_limits.get("main", ())]


def test_csv_logical_records(tmp_path):
    data = b'\xef\xbb\xbfname,amount\r\n"a, ""quote""\nline",\r\nb,0\r\nc,999\r\n'
    spec, man = case(tmp_path, data)
    bundle = load(spec, man, tmp_path)
    assert bundle.tables["main"].headers == ("name", "amount")
    assert len(bundle.tables["main"].rows) == 2
    assert [(c.kind, c.value, c.coordinate) for c in bundle.tables["main"].rows[0]] == [
        ("text", 'a, "quote"\nline', "A2"),
        ("blank", None, "B2"),
    ]
    assert bundle.tables["main"].rows[1][1].value == "0"
    assert bundle.snapshots["work"] == data
    (tmp_path / "input.csv").write_bytes(b"changed")
    assert bundle.snapshots["work"] == data
    assert not bundle.business_limits


@pytest.mark.parametrize(
    "data",
    [
        b'name,amount\n"unclosed,1\n',
        b'name,amount\na"bad,1\n',
        b'name,amount\n"bad"tail,1\n',
        b"name,amount\na,1,2\n",
        b"name,amount\na,\xff\n",
    ],
)
def test_invalid_csv_is_case_error(tmp_path, data):
    spec, man = case(tmp_path, data)
    with pytest.raises(CaseError):
        load(spec, man, tmp_path)


def test_xlsx_types_cache_and_false_dimension(tmp_path):
    data = workbook_bytes(
        ("amount", "flag", "day", "error", "formula"),
        (
            (0, True, datetime(2024, 1, 2), "#DIV/0!", "=1+1"),
            (2, False, datetime(2024, 1, 3), "#N/A", "=3+3"),
        ),
    )
    data = rewrite_zip(
        data,
        {
            "xl/worksheets/sheet1.xml": lambda raw: re.sub(
                rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', raw
            ).replace(b"<f>1+1</f><v /></c>", b"<f>1+1</f><v>2</v></c>")
        },
    )
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert b"<f>1+1</f><v>2</v>" in archive.read("xl/worksheets/sheet1.xml")
    cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        # A deliberately false dimension cannot prevent explicit cache probing either.
        assert (
            next(cached["Main"].iter_rows(min_row=2, max_row=2, min_col=5, max_col=5))[0].value == 2
        )
    finally:
        cached.close()
    spec, man = case(tmp_path, data, format="xlsx")
    bundle = load(spec, man, tmp_path)
    row = bundle.tables["main"].rows[0]
    assert [(c.kind, c.value) for c in row] == [
        ("number", "0"),
        ("bool", True),
        ("date", "2024-01-02T00:00:00"),
        ("error", "#DIV/0!"),
        ("formula", "=1+1"),
    ]
    assert len(bundle.tables["main"].rows) == 2
    (tmp_path / "input.xlsx").write_bytes(b"broken")
    assert bundle.snapshots["work"] == data


@pytest.mark.parametrize("budget", ["members", "member_bytes", "total_bytes"])
def test_xlsx_actual_zip_budgets_are_business_limits(tmp_path, budget):
    base = workbook_bytes()
    with zipfile.ZipFile(io.BytesIO(base)) as zf:
        count = len(zf.infolist())
    if budget == "members":
        extra = [(f"padding/{n}", b"") for n in range(201 - count)]
    elif budget == "member_bytes":
        extra = [("padding/big", b"x" * (16 * 1024 * 1024 + 1))]
    else:
        extra = [(f"padding/{n}", b"x" * (16 * 1024 * 1024)) for n in range(4)]
        extra.append(("padding/end", b"x"))
    data = rewrite_zip(base, extra=extra)
    spec, man = case(tmp_path, data, format="xlsx")
    bundle = load(spec, man, tmp_path)
    assert reasons(bundle) == ["UNSUPPORTED_INPUT_LIMIT"]
    assert "main" not in bundle.tables
    assert bundle.snapshots["work"] == data


@pytest.mark.parametrize("mutation", ["bad_zip", "duplicate", "dtd", "entity"])
def test_unsafe_xlsx_is_case_error(tmp_path, mutation):
    data = workbook_bytes()
    if mutation == "bad_zip":
        data = b"PK\x03\x04broken"
    elif mutation == "duplicate":
        with pytest.warns(UserWarning):
            data = rewrite_zip(data, extra=(("xl/workbook.xml", b"duplicate"),))
    else:
        declaration = (
            b'<!DOCTYPE worksheet [<!ENTITY x "boom">]>'
            if mutation == "entity"
            else b"<!DOCTYPE worksheet>"
        )
        data = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda raw: declaration + raw})
    spec, man = case(tmp_path, data, format="xlsx")
    with pytest.raises(CaseError):
        load(spec, man, tmp_path)


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_header_order_missing_ambiguity_and_width(tmp_path, format):
    def run(headers, bounds=()):
        data = (
            (",".join(headers) + "\n" + ",".join("1" for _ in headers) + "\n").encode()
            if format == "csv"
            else workbook_bytes(headers, (tuple(1 for _ in headers),))
        )
        spec, man = case(tmp_path, data, format=format, rows=(2, 2), boundaries=bounds)
        return load(spec, man, tmp_path)

    assert run(("amount", "name")).tables["main"].headers == ("amount", "name")
    with pytest.raises(CaseError):
        run(("name", "Amount"))
    assert reasons(run(("name", "Amount"), ("MISSING_FIELD",))) == ["MISSING_FIELD"]
    with pytest.raises(CaseError):
        run(("amount", "amount"))
    ambiguous = run(("amount", "amount"), ("AMBIGUOUS_FIELD",))
    assert ambiguous.tables["main"].headers == ("amount", "amount")
    assert reasons(ambiguous) == ["AMBIGUOUS_FIELD"]
    wide = run(tuple(f"column{n}" for n in range(100)) + ("amount",))
    assert reasons(wide) == ["UNSUPPORTED_INPUT_LIMIT"]
    assert "main" not in wide.tables


def test_count_keeps_rows_without_amount_field(tmp_path):
    spec, man = case(tmp_path, b"name,other\na,\nb,\n", metric={"op": "count", "source": "main"})
    bundle = load(spec, man, tmp_path)
    assert len(bundle.tables["main"].rows) == 2
    assert not bundle.business_limits


def test_missing_files_hash_and_root_containment(tmp_path):
    spec, man = case(tmp_path, b"name,amount\na,1\nb,2\n")
    entry = man.inputs[0]
    (tmp_path / "input.csv").unlink()
    with pytest.raises(CaseError):
        load(spec, man, tmp_path)
    missing = dataclasses.replace(
        man,
        inputs=(dataclasses.replace(entry, expected_missing=True, sha256=None),),
        intentional_boundaries=("MISSING_SOURCE",),
    )
    assert reasons(load(spec, missing, tmp_path)) == ["MISSING_SOURCE"]
    with pytest.raises(CaseError):
        load(spec, dataclasses.replace(missing, intentional_boundaries=()), tmp_path)
    (tmp_path / "input.csv").write_bytes(b"changed")
    with pytest.raises(CaseError):
        load(spec, man, tmp_path)
    with pytest.raises(CaseError):
        load(spec, missing, tmp_path)
    outside = tmp_path.parent / (tmp_path.name + "-outside.csv")
    outside.write_bytes(b"name,amount\na,1\nb,2\n")
    (tmp_path / "input.csv").unlink()
    (tmp_path / "input.csv").symlink_to(outside)
    with pytest.raises(CaseError):
        load(spec, man, tmp_path)


def test_range_limit_is_business_abstain(tmp_path):
    spec, man = case(tmp_path, b"name,amount\na,1\n", rows=(2, 1002))
    assert reasons(load(spec, man, tmp_path)) == ["UNSUPPORTED_INPUT_LIMIT"]


def test_validate_case_reads_stable_task_and_confirmation(tmp_path):
    spec, man = case(tmp_path, b"name,amount\na,1\nb,2\n")
    task = {"description": "Sum amount from main rows 2..3 in CNY, HALF_UP 2 places."}
    (tmp_path / "task.json").write_text(json.dumps(task), encoding="utf-8")
    from sheetbenchkit.artifacts import canonical_json, sha256_bytes

    man = dataclasses.replace(
        man,
        business_confirmation={
            "reviewed_contract_sha256": sha256_bytes(canonical_json(spec)),
            "task_sha256": sha256_bytes(canonical_json(task)),
        },
    )
    validate = getattr(contracts, "validate_case", None)
    assert callable(validate), "Missing T2 validate_case"
    validated = validate(spec, man, tmp_path)
    assert validated.task == task
    assert len(validated.inputs.tables["main"].rows) == 2
    (tmp_path / "task.json").write_text('{"description":"tampered"}')
    with pytest.raises(CaseError):
        validate(spec, man, tmp_path)


def test_over_file_limit_does_not_claim_hash_or_freeze_snapshot(tmp_path):
    from sheetbenchkit.artifacts import canonical_json, freeze_suite, sha256_bytes
    from sheetbenchkit.models import CandidateCase, ExpectedResult, ValidatedCase

    data = b"name,amount\na," + b"1" * (10 * 1024 * 1024) + b"\n"
    spec, man = case(tmp_path, data, rows=(2, 2))
    bundle = load(spec, man, tmp_path)
    issue = bundle.business_limits["main"][0]
    assert issue["reason"] == "UNSUPPORTED_INPUT_LIMIT"
    assert issue["details"]["snapshot_unavailable"] is True
    assert issue["details"]["hash_verified"] is False
    assert "work" not in bundle.snapshots
    assert "main" not in bundle.tables
    task = {"description": "Sum declared amount."}
    confirmation = {
        "reviewed_contract_sha256": sha256_bytes(canonical_json(spec)),
        "task_sha256": sha256_bytes(canonical_json(task)),
    }
    man = dataclasses.replace(man, business_confirmation=confirmation)
    candidate = CandidateCase(
        ValidatedCase(spec, man, bundle, task),
        ExpectedResult("ABSTAIN", None, None, "UNSUPPORTED_INPUT_LIMIT", ()),
        "preset_independent",
        confirmation,
        {},
    )
    destination = tmp_path / "frozen"
    with pytest.raises(CaseError, match="INPUT_HASH_MISMATCH"):
        freeze_suite((candidate,), (), destination)
    assert not destination.exists()


@pytest.mark.parametrize("data,rows", [(b"name,amount\na,1\n", (2, 3)), (b"", (2, 2))])
def test_csv_nonexistent_logical_records_are_error(tmp_path, data, rows):
    spec, man = case(tmp_path, data, rows=rows, boundaries=("MISSING_SOURCE",))
    with pytest.raises(CaseError, match="INVALID_CSV_RANGE"):
        load(spec, man, tmp_path)


def test_distractor_zip_limit_is_recorded_by_file_id(tmp_path):
    spec, man = case(tmp_path, b"name,amount\na,1\nb,2\n")
    extra = rewrite_zip(workbook_bytes(), extra=(("padding/big", b"x" * (16 * 1024 * 1024 + 1)),))
    path = tmp_path / "distractor.xlsx"
    path.write_bytes(extra)
    entry = dataclasses.replace(
        man.inputs[0],
        file_id="distractor",
        relative_path=path.name,
        format="xlsx",
        sha256=hashlib.sha256(extra).hexdigest(),
    )
    man = dataclasses.replace(man, inputs=man.inputs + (entry,))
    bundle = load(spec, man, tmp_path)
    assert bundle.business_limits["file:distractor"][0]["reason"] == "UNSUPPORTED_INPUT_LIMIT"
    assert bundle.snapshots["distractor"] == extra
    assert len(bundle.tables["main"].rows) == 2


def test_xlsx_sparse_declared_rows_are_blank_without_trusting_dimension(tmp_path):
    data = workbook_bytes(rows=(("a", 0),))
    spec, man = case(tmp_path, data, format="xlsx", rows=(4, 5))
    bundle = load(spec, man, tmp_path)
    assert len(bundle.tables["main"].rows) == 2
    assert [(cell.kind, cell.coordinate) for row in bundle.tables["main"].rows for cell in row] == [
        ("blank", "A4"),
        ("blank", "B4"),
        ("blank", "A5"),
        ("blank", "B5"),
    ]


def test_external_links_and_formula_are_inert(tmp_path):
    data = workbook_bytes(rows=(("a", "='[1]Book.xlsx'!A1"),))
    data = rewrite_zip(
        data,
        {
            "xl/workbook.xml": lambda raw: raw.replace(
                b"</workbook>",
                b"<externalReferences>"
                b'<externalReference xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                b'relationships" r:id="external1"/></externalReferences></workbook>',
            ),
            "xl/_rels/workbook.xml.rels": lambda raw: raw.replace(
                b"</Relationships>",
                b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                b'relationships/externalLink" Target="externalLinks/externalLink1.xml" '
                b'Id="external1"/></Relationships>',
            ),
        },
        extra=(
            (
                "xl/externalLinks/externalLink1.xml",
                b'<externalLink xmlns="http://schemas.openxmlformats.'
                b'org/spreadsheetml/2006/main"><externalBook xmlns:r="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships" r:id="rId1"/></externalLink>',
            ),
            (
                "xl/externalLinks/_rels/externalLink1.xml.rels",
                b'<Relationships xmlns="http://schemas.'
                b'openxmlformats.org/package/2006/relationships"><Relationship '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                b'externalLinkPath" '
                b'Target="https://invalid.example/never-fetch.xlsx" TargetMode="External" '
                b'Id="rId1"/></Relationships>',
            ),
        ),
    )
    spec, man = case(tmp_path, data, format="xlsx", rows=(2, 2))
    cell = load(spec, man, tmp_path).tables["main"].rows[0][1]
    assert (cell.kind, cell.value) == ("formula", "='[1]Book.xlsx'!A1")


def test_xlsx_data_table_formula_preserves_formula_type(tmp_path):
    data = workbook_bytes(rows=(("a", "=1+1"),))
    data = rewrite_zip(
        data,
        {
            "xl/worksheets/sheet1.xml": lambda raw: raw.replace(
                b"<f>1+1</f><v />", b'<f t="dataTable" ref="B2:B2" r1="A2"/><v>42</v>'
            )
        },
    )
    spec, man = case(tmp_path, data, format="xlsx", rows=(2, 2))
    cell = load(spec, man, tmp_path).tables["main"].rows[0][1]
    assert cell.kind == "formula"
    assert cell.value == {"t": "dataTable", "ref": "B2:B2", "r1": "A2"}


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_declared_exact_range_and_filters_keep_all_rows(tmp_path, format):
    from sheetbenchkit.models import FilterSpec

    data = (
        b"name,amount\na,1\nb,2\nc,999\n"
        if format == "csv"
        else workbook_bytes(rows=(("a", 1), ("b", 2), ("c", 999)))
    )
    spec, man = case(tmp_path, data, format=format)
    spec = dataclasses.replace(
        spec, metric=dataclasses.replace(spec.metric, filter=FilterSpec("name", "a"))
    )
    bundle = load(spec, man, tmp_path)
    assert len(bundle.tables["main"].rows) == 2
    assert [row[0].value for row in bundle.tables["main"].rows] == ["a", "b"]
    assert [row[1].value for row in bundle.tables["main"].rows] == ["1", "2"]
    spec = dataclasses.replace(
        spec, metric=dataclasses.replace(spec.metric, filter=FilterSpec("missing_filter", "a"))
    )
    with pytest.raises(CaseError, match="UNDECLARED_BOUNDARY"):
        load(spec, man, tmp_path)
    rejected = load(
        spec, dataclasses.replace(man, intentional_boundaries=("MISSING_FIELD",)), tmp_path
    )
    assert reasons(rejected) == ["MISSING_FIELD"]


def test_read_only_workbook_archive_is_closed_after_load(tmp_path, monkeypatch):
    module = importlib.import_module("sheetbenchkit.loaders")
    opened = []
    original = module.load_workbook

    def track(*args, **kwargs):
        workbook = original(*args, **kwargs)
        opened.append(workbook)
        return workbook

    monkeypatch.setattr(module, "load_workbook", track)
    spec, man = case(tmp_path, workbook_bytes(), format="xlsx")
    bundle = load(spec, man, tmp_path)
    assert bundle.tables["main"].rows[0][1].value == "0"
    assert len(opened) == 1
    assert opened[0]._archive.fp is None


def test_nontext_xlsx_header_is_explicit_business_limit(tmp_path):
    spec, man = case(tmp_path, workbook_bytes(("name", 7)), format="xlsx")
    bundle = load(spec, man, tmp_path)
    assert reasons(bundle) == ["UNSUPPORTED_CELL"]
    assert "main" not in bundle.tables


def test_missing_xlsx_sheet_requires_declared_boundary(tmp_path):
    spec, man = case(tmp_path, workbook_bytes(), format="xlsx", sheet="Absent")
    with pytest.raises(CaseError, match="UNDECLARED_BOUNDARY"):
        load(spec, man, tmp_path)
    assert reasons(
        load(spec, dataclasses.replace(man, intentional_boundaries=("MISSING_SOURCE",)), tmp_path)
    ) == ["MISSING_SOURCE"]


def test_csv_exact_file_and_table_limits_are_supported(tmp_path):
    prefix = b"name,amount\n"
    data = prefix + b"a" * (10 * 1024 * 1024 - len(prefix) - 3) + b",0\n"
    spec, man = case(tmp_path, data, rows=(2, 2))
    assert load(spec, man, tmp_path).tables["main"].rows[0][1].value == "0"
    headers = ("amount",) + tuple(f"column{n}" for n in range(99))
    data = (",".join(headers) + "\n" + (",".join("0" for _ in headers) + "\n") * 1000).encode()
    spec, man = case(tmp_path, data, rows=(2, 1001))
    table = load(spec, man, tmp_path).tables["main"]
    assert len(table.headers) == 100
    assert len(table.rows) == 1000
    assert len(table.rows[-1]) == 100
    assert table.rows[-1][-1].coordinate == "CV1001"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_source_named_file_keeps_its_limit_binding(tmp_path, format):
    headers = tuple(f"column{n}" for n in range(100)) + ("amount",)
    data = (
        (",".join(headers) + "\n" + ",".join("1" for _ in headers) + "\n").encode()
        if format == "csv"
        else workbook_bytes(headers, (tuple(1 for _ in headers),))
    )
    spec, man = case(tmp_path, data, format=format, rows=(2, 2))
    spec = dataclasses.replace(
        spec,
        sources={"file": spec.sources["main"]},
        metric=dataclasses.replace(spec.metric, source="file"),
    )
    bundle = load(spec, man, tmp_path)
    assert bundle.business_limits["file"][0]["reason"] == "UNSUPPORTED_INPUT_LIMIT"
    assert "file:work" not in bundle.business_limits


def test_nonregular_input_is_case_error_without_waiting_for_a_writer(tmp_path):
    import os
    import signal

    if not hasattr(os, "mkfifo") or not hasattr(signal, "SIGALRM"):
        pytest.skip("FIFO/interrupt probe requires Unix")
    spec, man = case(tmp_path, b"name,amount\na,1\nb,2\n")
    path = tmp_path / "input.csv"
    path.unlink()
    os.mkfifo(path)

    timeouts = []

    def timeout(signum, frame):
        timeouts.append(True)
        raise TimeoutError("loader blocked opening a nonregular input")

    previous = signal.signal(signal.SIGALRM, timeout)
    signal.setitimer(signal.ITIMER_REAL, 0.2)
    try:
        with pytest.raises(CaseError):
            load(spec, man, tmp_path)
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)
    assert not timeouts, "loader waited for a writer before rejecting the FIFO"


def test_corrupt_lzma_member_is_case_error_at_public_loader(tmp_path):
    import struct

    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_LZMA) as archive:
        archive.writestr("padding/lzma", b"x" * 100)
    corrupt = bytearray(stream.getvalue())
    name_size, extra_size = struct.unpack_from("<HH", corrupt, 26)
    # ZIP_LZMA payload: version(2), property-size(2), then LZMA properties.
    corrupt[30 + name_size + extra_size + 4] = 255
    spec, man = case(tmp_path, bytes(corrupt), format="xlsx")
    with pytest.raises(CaseError) as caught:
        load(spec, man, tmp_path)
    assert caught.value.reason == "INVALID_XLSX"
    assert caught.value.details == "LZMAError"
