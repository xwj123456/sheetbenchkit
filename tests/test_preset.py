"""The approved fixed matrix must catch business mistakes, independently of gold creation."""

import copy
import csv
import importlib
import io
import json
import subprocess
import sys
from dataclasses import replace
from datetime import datetime
from importlib.resources import files
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from sheetbenchkit import artifacts as a
from sheetbenchkit import models as m
from sheetbenchkit.cli import main
from sheetbenchkit.contracts import decode_document, to_document, validate_schema
from sheetbenchkit.grader import grade_case, grade_suite
from sheetbenchkit.runner import make_task_envelope

VALUES = {
    "S01": "100",
    "S02": "137",
    "S03": "100",
    "S04": "100",
    "S05": "120",
    "S06": "100",
    "S07": "120",
    "S08": "180",
    "S09": "120",
    "S10": "2",
    "S11": "1",
    "S12": "2",
    "N13": "5",
    "N14": "0",
    "N16": "0",
    "N18": "0.0313",
    "B20": "30",
    "B21": "30",
    "R23": "0.8273",
    "R24": "0.7500",
    "R25": "0.0000",
    "U27": "30",
    "U28": "0.5000",
}
REASONS = {
    "N15": "MISSING_VALUE",
    "N17": "INVALID_VALUE",
    "B19": "AMBIGUOUS_FIELD",
    "B22": "MISSING_FIELD",
    "R26": "ZERO_DENOMINATOR",
    "U29": "UNSUPPORTED_UNIT_CONVERSION",
    "U30": "UNSUPPORTED_UNIT_CONVERSION",
}
KILLERS = {
    "M01": ("S02", "S03"),
    "M02": ("S05", "S06"),
    "M03_contains": ("S08", "S09", "B20"),
    "M03_index": ("B21",),
    "M04": ("S11", "S12"),
    "M05": ("N15",),
    "M06": ("N17",),
    "M07": ("R23",),
    "M08_scale": ("U28",),
    "M08_unit": ("U29", "U30"),
    "M09": ("N18",),
    "M10": tuple(VALUES),
    "M11": ("S01", "B22", "R26"),
    "M12": ("S01", "S04", "S07"),
}


def preset():
    root = Path(str(files("sheetbenchkit").joinpath("data/v1")))
    assert (root / "suite.json").is_file(), "Missing Task8 frozen 30-case preset"
    suite = decode_document("suite", (root / "suite.json").read_bytes())
    return root, suite, a.preflight_suite(suite, root)


def adapter():
    try:
        module = importlib.import_module("sheetbenchkit.examples.independent_runner")
    except ModuleNotFoundError:
        pytest.fail("Missing Task8 independent file-reading adapter")
    assert callable(getattr(module, "solve", None)), "Missing independent solve API"
    return module


def envelope(case, root):
    result = make_task_envelope(case)
    return to_document(
        replace(
            result,
            inputs=tuple(
                replace(
                    entry,
                    relative_path=str(
                        (root / "cases" / case.gold.case_id / entry.relative_path).resolve()
                    ),
                )
                for entry in result.inputs
            ),
        )
    )


def observation(case, result, config="correct"):
    return m.Observation(
        case.gold.case_id,
        1,
        config,
        "SUCCESS",
        json.dumps(result, ensure_ascii=False),
        "",
        None,
        {},
    )


def test_frozen_30_matrix_and_independent_correct_adapter():
    root, suite, cases = preset()
    assert len(cases) == 30 and len({c.gold.case_id for c in cases}) == 30
    assert {c.gold.case_id for c in cases} == set(VALUES) | set(REASONS)
    assert len(suite.families) == 4
    module = adapter()
    observations = []
    for case in cases:
        want = case.gold.expected
        identity = case.gold.case_id
        assert case.gold.origin == "preset_independent"
        assert case.gold.confirmation["review_method"] == "AI independent design review"
        if identity in VALUES:
            assert (want.status, want.value, want.reason) == ("VALUE", VALUES[identity], None)
            unit = (
                "count"
                if identity in ("S10", "S11", "S12")
                else "ratio"
                if identity in ("N18", "R23", "R24", "R25", "U28")
                else "CNY"
                if identity == "U27"
                else "units"
            )
            assert want.unit == unit
        else:
            assert (want.status, want.value, want.unit, want.reason) == (
                "ABSTAIN",
                None,
                None,
                REASONS[identity],
            )
            assert want.reason in case.validated.manifest.intentional_boundaries
        output = module.solve(envelope(case, root))
        validate_schema("result", output)
        observations.append(observation(case, output))
    report = grade_suite(
        suite, cases, m.ObservationBatch("1", (m.RunSpec("correct", 1),), tuple(observations))
    )
    assert report.valid and report.status_counts["PASS"] == 30
    assert report.attempted == report.observed_count == 30 and report.missing_count == 0
    assert len(report.family_grades) == 4
    assert all(item.verdict == "PASS" for item in report.family_grades)


@pytest.mark.parametrize("control,killers", KILLERS.items())
def test_each_negative_control_has_legal_business_killers(control, killers):
    root, _, cases = preset()
    try:
        mutants = importlib.import_module("controls.mutants")
    except ModuleNotFoundError:
        pytest.fail("Missing Task8 fourteen independent negative-control subtypes")
    assert set(mutants.CONTROLS) == set(KILLERS)
    by_id = {case.gold.case_id: case for case in cases}
    for identity in killers:
        case = by_id[identity]
        payload = envelope(case, root)
        output = mutants.solve(control, payload)
        validate_schema("result", output)
        correct_bindings = adapter().declared_bindings(payload["output_protocol"]["contract"])
        if control != "M12":
            assert output["bindings"] == correct_bindings
        grade = grade_case(case, observation(case, output))
        assert grade.verdict == "FAIL", (control, identity, grade)
        if control == "M12":
            assert grade.value_check["verdict"] == "PASS"
            assert grade.binding_check["verdict"] == "FAIL"
        else:
            assert grade.value_check["verdict"] == "FAIL"
            assert grade.binding_check["verdict"] == (
                "PASS" if case.gold.expected.status == "VALUE" else "NOT_CHECKED"
            )


def test_family_files_change_only_one_declared_cell():
    root, suite, cases = preset()
    by_id = {case.gold.case_id: case for case in cases}

    def cells(case):
        result = {}
        for entry in case.validated.manifest.inputs:
            data = (root / "cases" / case.gold.case_id / entry.relative_path).read_bytes()
            if entry.format == "xlsx":
                book = load_workbook(io.BytesIO(data), data_only=False)
                try:
                    for sheet in book:
                        for row in sheet:
                            for cell in row:
                                result[(entry.file_id, sheet.title, cell.coordinate)] = cell.value
                finally:
                    book.close()
        return result

    targets = {"file": "137", "sheet": "120", "field": "180", "filter-count": "1"}
    for family in suite.families:
        base = by_id[family.base]
        assert len(base.validated.manifest.variant_recipes) == 2
        assert by_id[family.target].gold.expected.value == targets[family.family_id]
        for identity, role in ((family.target, "target"), (family.distractor, "distractor")):
            changed = by_id[identity]
            assert changed.validated.spec == base.validated.spec
            assert changed.validated.task == base.validated.task
            assert not changed.validated.manifest.variant_recipes
            before, after = cells(base), cells(changed)
            assert before.keys() == after.keys()
            differences = [key for key in before if before[key] != after[key]]
            assert len(differences) == 1
            location = family.mutation_locations[role]
            assert differences == [(location["file_id"], location["sheet"], location["coordinate"])]
    family_members = {
        identity for f in suite.families for identity in (f.base, f.target, f.distractor)
    }
    assert sum(c.gold.case_id not in family_members for c in cases) == 18


def test_demo_generation_selects_22_bases_and_preserves_frozen_bytes(tmp_path):
    root, _, _ = preset()
    before = {
        p.relative_to(root): a.sha256_bytes(p.read_bytes()) for p in root.rglob("*") if p.is_file()
    }
    destination = tmp_path / "generated"
    assert main(["generate", "--demo", "--seed", "7", "--output", str(destination)]) == 0
    generated = decode_document("suite", (destination / "suite.json").read_bytes())
    cases = a.preflight_suite(generated, destination)
    assert len(cases) == 30 and len(generated.families) == 4
    assert sum(c.gold.expected.status == "VALUE" for c in cases) == 23
    assert all(c.gold.origin == "generated_contract" for c in cases)
    by_id = {case.gold.case_id: case for case in cases}
    triples = {
        "S01": ["100", "137", "100"],
        "S04": ["100", "120", "100"],
        "S07": ["120", "180", "120"],
        "S10": ["2", "1", "2"],
    }
    for family in generated.families:
        base = by_id[family.base].gold.evidence["base_case_id"]
        assert [
            by_id[identity].gold.expected.value
            for identity in (family.base, family.target, family.distractor)
        ] == triples[base]
    module = adapter()
    observations = tuple(
        observation(case, module.solve(envelope(case, destination))) for case in cases
    )
    report = grade_suite(
        generated, cases, m.ObservationBatch("1", (m.RunSpec("correct", 1),), observations)
    )
    assert report.status_counts["PASS"] == 30
    assert all(item.verdict == "PASS" for item in report.family_grades)
    assert before == {
        p.relative_to(root): a.sha256_bytes(p.read_bytes()) for p in root.rglob("*") if p.is_file()
    }


def test_adapter_reads_absolute_files_and_does_not_lookup_case_outputs(tmp_path):
    root, _, cases = preset()
    case = next(c for c in cases if c.gold.case_id == "N13")
    payload = copy.deepcopy(envelope(case, root))
    entry = payload["inputs"][0]
    physical = tmp_path / "fresh 中文 data.csv"
    physical.write_text("amount\n7\n8\n", encoding="utf-8")
    entry["relative_path"] = str(physical.resolve())
    payload["case_id"] = "never_seen_before"
    module = adapter()
    assert module.solve(payload)["value"] == "15"
    completed = subprocess.run(
        [sys.executable, "-m", module.__name__],
        input=json.dumps(payload),
        text=True,
        capture_output=True,
        cwd=tmp_path,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout)["value"] == "15"
    physical.write_text("amount\n0\n-8\n", encoding="utf-8")
    assert module.solve(payload)["value"] == "-8"
    entry["relative_path"] = "relative.csv"
    with pytest.raises(ValueError, match="absolute"):
        module.solve(payload)


def boundary_payload(tmp_path, raw, *, format="csv", kind=None, filtering=None, places=0):
    """Real external input under the approved one-source contract, with unseen identity."""
    root, _, cases = preset()
    payload = copy.deepcopy(envelope(next(c for c in cases if c.gold.case_id == "N13"), root))
    payload["case_id"] = "custom_boundary"
    spec = payload["output_protocol"]["contract"]
    entry = payload["inputs"][0]
    field = "status" if filtering is not None else "amount"
    if format == "csv":
        physical = tmp_path / "boundary.csv"
        with physical.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows([[field], [raw]])
    else:
        physical = tmp_path / "boundary.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Actual"
        sheet.append([field])
        sheet.cell(2, 1, raw)
        if kind is not None:
            sheet.cell(2, 1).data_type = kind
        book.save(physical)
        book.close()
        spec["sources"]["chosen"]["sheet"] = "Actual"
        payload["output_protocol"]["unit_labels"] = {"data": {"Actual": {"amount": "units"}}}
    entry["format"] = format
    entry["relative_path"] = str(physical.resolve())
    spec["sources"]["chosen"]["rows"] = [2, 2]
    if filtering is not None:
        spec["metric"] = {
            "op": "count",
            "source": "chosen",
            "filter": {"field": "status", "eq": filtering},
        }
    else:
        spec["metric"]["places"] = places
    return payload


def assert_boundary_result(payload, expected):
    result = adapter().solve(payload)
    validate_schema("result", result)
    assert (result["status"], result["value"], result["unit"], result["reason"]) == expected
    assert result["bindings"] == adapter().declared_bindings(payload["output_protocol"]["contract"])


@pytest.mark.parametrize(
    "raw,status,value,reason",
    [
        ("0" * 79 + "1", "VALUE", "1.000000000000", None),
        ("0" * 80 + "1", "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("9" * 30, "VALUE", "999999999999999999999999999999.000000000000", None),
        ("9" * 31, "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("0" * 40 + "9" * 30, "VALUE", "999999999999999999999999999999.000000000000", None),
        ("0" * 40 + "9" * 31, "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("0." + "0" * 29 + "1", "VALUE", "0.000000000000", None),
        ("0." + "0" * 30 + "1", "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("1e12", "VALUE", "1000000000000.000000000000", None),
        ("1e13", "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("1e-12", "VALUE", "0.000000000001", None),
        ("1e-13", "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
        ("", "ABSTAIN", None, "MISSING_VALUE"),
        ("   ", "ABSTAIN", None, "MISSING_VALUE"),
        (" \t ", "ABSTAIN", None, "MISSING_VALUE"),
        (" 1", "ABSTAIN", None, "INVALID_VALUE"),
        ("1 ", "ABSTAIN", None, "INVALID_VALUE"),
        ("oops", "ABSTAIN", None, "INVALID_VALUE"),
        ("0" * 81 + "oops", "ABSTAIN", None, "INVALID_VALUE"),
        ("=1+1", "ABSTAIN", None, "INVALID_VALUE"),
        ("#DIV/0!", "ABSTAIN", None, "INVALID_VALUE"),
    ],
)
def test_independent_csv_numeric_limit_and_missing_reasons(tmp_path, raw, status, value, reason):
    payload = boundary_payload(tmp_path, raw, places=12)
    assert_boundary_result(payload, (status, value, "units" if status == "VALUE" else None, reason))


@pytest.mark.parametrize(
    "raw,eq,value",
    [
        ("=OK", "=OK", "1"),
        ("#DIV/0!", "#DIV/0!", "1"),
        ("1", "1", "1"),
        ("YES", "YES", "1"),
        ("yes", "YES", "0"),
        ("", "YES", "0"),
        ("   ", "   ", "0"),
        ("\t", "\t", "0"),
        (" YES", "YES", "0"),
    ],
)
def test_independent_csv_exact_literal_filter_and_blank_text(tmp_path, raw, eq, value):
    payload = boundary_payload(tmp_path, raw, filtering=eq)
    assert_boundary_result(payload, ("VALUE", value, "count", None))


@pytest.mark.parametrize(
    "raw,kind,status,value,reason",
    [
        ("=1+1", "f", "ABSTAIN", None, "UNSUPPORTED_CELL"),
        ("=1+1", "s", "ABSTAIN", None, "INVALID_VALUE"),
        ("#DIV/0!", "e", "ABSTAIN", None, "UNSUPPORTED_CELL"),
        ("#DIV/0!", "s", "ABSTAIN", None, "INVALID_VALUE"),
        (True, "b", "ABSTAIN", None, "UNSUPPORTED_CELL"),
        (datetime(2026, 10, 2), None, "ABSTAIN", None, "UNSUPPORTED_CELL"),
        (7, "n", "VALUE", "7", None),
        ("7", "s", "VALUE", "7", None),
        (None, None, "ABSTAIN", None, "MISSING_VALUE"),
        ("   ", "s", "ABSTAIN", None, "MISSING_VALUE"),
        ("1e13", "s", "ABSTAIN", None, "UNSUPPORTED_INPUT_LIMIT"),
    ],
)
def test_independent_xlsx_numeric_preserves_intrinsic_kinds(
    tmp_path, raw, kind, status, value, reason
):
    payload = boundary_payload(tmp_path, raw, format="xlsx", kind=kind)
    assert_boundary_result(payload, (status, value, "units" if status == "VALUE" else None, reason))


@pytest.mark.parametrize(
    "raw,kind,eq,status,value,reason",
    [
        ("=OK", "f", "=OK", "ABSTAIN", None, "INVALID_FILTER_VALUE"),
        ("=OK", "s", "=OK", "VALUE", "1", None),
        ("#DIV/0!", "e", "#DIV/0!", "ABSTAIN", None, "INVALID_FILTER_VALUE"),
        ("#DIV/0!", "s", "#DIV/0!", "VALUE", "1", None),
        (1, "n", "1", "ABSTAIN", None, "INVALID_FILTER_VALUE"),
        (True, "b", "True", "ABSTAIN", None, "INVALID_FILTER_VALUE"),
        (datetime(2026, 10, 2), None, "2026-10-02", "ABSTAIN", None, "INVALID_FILTER_VALUE"),
        ("2026-10-02", "s", "2026-10-02", "VALUE", "1", None),
        ("1", "s", "1", "VALUE", "1", None),
        ("YES", "s", "YES", "VALUE", "1", None),
        (None, None, "YES", "VALUE", "0", None),
        ("", "s", "YES", "VALUE", "0", None),
        ("   ", "s", "   ", "VALUE", "0", None),
    ],
)
def test_independent_xlsx_filter_preserves_text_formula_error_and_nontext(
    tmp_path, raw, kind, eq, status, value, reason
):
    payload = boundary_payload(tmp_path, raw, format="xlsx", kind=kind, filtering=eq)
    assert_boundary_result(payload, (status, value, "count" if status == "VALUE" else None, reason))
