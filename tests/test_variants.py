"""Real frozen-file checks catch copied, invisible, or wrongly bound mutations."""

from __future__ import annotations

import importlib
import importlib.util
import io
import re
import zipfile
from dataclasses import replace
from decimal import Inexact, localcontext
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook

from sheetbenchkit.artifacts import canonical_json, sha256_bytes
from sheetbenchkit.contracts import decode_document, validate_case
from sheetbenchkit.models import (
    AggregateSumSpec,
    CaseError,
    CaseManifest,
    CountSpec,
    FilterSpec,
    FrozenCase,
    InputEntry,
    MutationRecipe,
    RatioSpec,
    RootSumSpec,
    RuleSpec,
    SourceSpec,
)


def generate(bases, seed, destination):
    assert importlib.util.find_spec("sheetbenchkit.variants"), "Missing T4 actual variants"
    return importlib.import_module("sheetbenchkit.variants").generate_suite(
        bases, seed, destination
    )


def recipe(kind, sheet, header=1, rows=(2, 3), row=2, field="amount", delta="37", text=None):
    return MutationRecipe(
        kind,
        "work",
        sheet,
        header,
        rows,
        field,
        row,
        delta if kind.endswith("numeric") else None,
        text if kind.endswith("filter") else None,
    )


def make_base(root: Path, format="csv", amount="100", places=0, filtering=False, shuffle=False):
    root.mkdir()
    if format == "csv":
        data = f"flag,amount,note\nyes,{amount},target\nno,863,distractor\n".encode()
        sheet = "@csv"
        target = recipe("target_numeric", sheet)
        distractor = recipe("distractor_numeric", sheet, row=3)
    else:
        book = Workbook()
        main = book.active
        main.title = "Chosen"
        main.append(("flag", "amount", "note"))
        main.append(("yes", amount, "target"))
        main.append(("no", "863", "excluded"))
        other = book.create_sheet("Distractor")
        other["A1"] = "Keep this preamble"
        other.append(("preamble",))
        other["A5"], other["B5"], other["C5"] = "note", "amount", "flag"
        other["A6"], other["B6"], other["C6"] = "decoy", "863", "no"
        book.create_sheet("Untouched")["A1"] = "=1+2"
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        data = stream.getvalue()
        sheet = "Chosen"
        target = recipe("target_numeric", sheet)
        distractor = recipe("distractor_numeric", "Distractor", 5, (6, 6), 6)
    source = SourceSpec("work", sheet, 1, (2, 3))
    spec = RuleSpec(
        "1",
        "custom_total",
        {"chosen": source},
        RootSumSpec(
            "sum",
            "chosen",
            "amount",
            "CNY",
            FilterSpec("flag", "yes"),
            places=places,
            rounding="HALF_UP",
        ),
    )
    task = {
        "description": "Sum amount from chosen where flag equals yes, CNY, HALF_UP.",
        "places": places,
    }
    if filtering:
        target = recipe("target_filter", sheet, row=3, field="flag", text="yes")
        distractor = replace(
            distractor,
            kind="distractor_filter",
            field="flag",
            numeric_delta=None,
            text_replacement="yes",
        )
        if format == "csv":
            # A separate nonmetric column is a legal declared distractor filter location.
            distractor = replace(distractor, field="note", text_replacement="yes")
    recipes = (target, distractor)
    if shuffle:
        recipes += (recipe("column_shuffle", sheet, field=None, row=None, delta=None),)
    path = root / ("input." + format)
    path.write_bytes(data)
    (root / "task.json").write_bytes(canonical_json(task))
    confirmation = {
        "reviewed_contract_sha256": sha256_bytes(canonical_json(spec)),
        "task_sha256": sha256_bytes(canonical_json(task)),
        "reviewer": "custom-test",
    }
    manifest = CaseManifest(
        "1",
        "bespoke_" + format,
        "custom",
        (InputEntry("work", path.name, format, sha256_bytes(data), False),),
        {"work": {sheet: {"amount": "CNY"}}},
        (),
        confirmation,
        recipes,
    )
    return validate_case(spec, manifest, root)


def frozen_by_id(suite, path):
    result = {}
    for ref in suite.cases:
        spec = decode_document("contract", (path / ref.contract_path).read_bytes())
        manifest = decode_document("manifest", (path / ref.manifest_path).read_bytes())
        gold = decode_document("gold", (path / ref.gold_path).read_bytes())
        validated = validate_case(spec, manifest, (path / ref.manifest_path).parent)
        result[ref.case_id] = FrozenCase(validated, gold, "unused-test-only")
    return result


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_target_distractor_triplet(tmp_path, format):
    base = make_base(tmp_path / "original", format)
    original = (tmp_path / "original" / ("input." + format)).read_bytes()
    first, second = tmp_path / "first", tmp_path / "second"
    suite = generate((base,), 7, first)
    again = generate((base,), 7, second)
    assert suite.content_sha256 == again.content_sha256
    assert {p.relative_to(first): p.read_bytes() for p in first.rglob("*") if p.is_file()} == {
        p.relative_to(second): p.read_bytes() for p in second.rglob("*") if p.is_file()
    }
    assert (tmp_path / "original" / ("input." + format)).read_bytes() == original
    cases = frozen_by_id(suite, first)
    (family,) = suite.families
    members = [cases[key] for key in (family.base, family.target, family.distractor)]
    assert [case.gold.expected.value for case in members] == ["100", "137", "100"]
    assert len({case.validated.manifest.case_id for case in members}) == 3
    assert len({case.validated.manifest.inputs[0].sha256 for case in members}) == 3
    assert len({ref.task_sha256 for ref in suite.cases}) == 1
    assert len({ref.contract_sha256 for ref in suite.cases}) == 1
    assert all(case.gold.origin == "generated_contract" for case in members)
    assert all(case.gold.confirmation == base.manifest.business_confirmation for case in members)
    assert members[0].validated.inputs.snapshots["work"] == original
    assert family.mutation_locations
    if format == "xlsx":
        for index, case in enumerate(members):
            book = load_workbook(
                io.BytesIO(case.validated.inputs.snapshots["work"]), data_only=False
            )
            assert str(book["Chosen"]["B2"].value) == ("137" if index == 1 else "100")
            assert str(book["Distractor"]["B6"].value) == ("900" if index == 2 else "863")
            assert book["Distractor"]["A1"].value == "Keep this preamble"
            assert book["Untouched"]["A1"].value == "=1+2"
            book.close()
    else:
        assert members[1].validated.inputs.snapshots["work"] == (
            b"flag,amount,note\nyes,137,target\nno,863,distractor\n"
        )
        assert members[2].validated.inputs.snapshots["work"] == (
            b"flag,amount,note\nyes,100,target\nno,900,distractor\n"
        )


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_explicit_filter_recipes_and_exact_header_shuffle(tmp_path, format):
    base = make_base(tmp_path / "original", format, filtering=True, shuffle=True)
    out = tmp_path / "generated"
    suite = generate((base,), 7, out)
    cases = frozen_by_id(suite, out)
    (family,) = suite.families
    assert [
        cases[key].gold.expected.value for key in (family.base, family.target, family.distractor)
    ] == ["100", "963", "100"]
    (shuffled,) = [
        case
        for key, case in cases.items()
        if key not in {family.base, family.target, family.distractor}
    ]
    assert shuffled.gold.expected.value == "100"
    assert shuffled.validated.inputs.tables["chosen"].headers != ("flag", "amount", "note")
    assert set(shuffled.validated.inputs.tables["chosen"].headers) == {"flag", "amount", "note"}
    assert shuffled.validated.spec == base.spec
    assert shuffled.validated.manifest.inputs[0].sha256 != base.manifest.inputs[0].sha256


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_invisible_and_invalid_mutations(tmp_path, format):
    invisible = make_base(tmp_path / "original", format, amount="1.00001", places=4)
    recipes = (
        replace(invisible.manifest.variant_recipes[0], numeric_delta="0.00001"),
        invisible.manifest.variant_recipes[1],
    )
    invisible = replace(invisible, manifest=replace(invisible.manifest, variant_recipes=recipes))
    destination = tmp_path / "invisible"
    with pytest.raises(CaseError):
        generate((invisible,), 7, destination)
    assert not destination.exists()
    base = make_base(tmp_path / "other", format)
    bad = replace(base.manifest.variant_recipes[0], kind="distractor_numeric")
    base = replace(
        base,
        manifest=replace(base.manifest, variant_recipes=(base.manifest.variant_recipes[0], bad)),
    )
    destination = tmp_path / "wrong-binding"
    with pytest.raises(CaseError):
        generate((base,), 7, destination)
    assert not destination.exists()


def test_unconfirmed_and_unverified_snapshots_cannot_generate(tmp_path):
    base = make_base(tmp_path / "original")
    invalids = (
        replace(base, manifest=replace(base.manifest, business_confirmation={})),
        replace(base, inputs=replace(base.inputs, snapshots={})),
        replace(base, inputs=replace(base.inputs, snapshots={"work": b"tampered"})),
        replace(
            base, inputs=replace(base.inputs, snapshots={"work": b"x" * (10 * 1024 * 1024 + 1)})
        ),
    )
    for index, invalid in enumerate(invalids):
        destination = tmp_path / str(index)
        with pytest.raises(CaseError):
            generate((invalid,), 7, destination)
        assert not destination.exists()


@pytest.mark.parametrize(
    "change",
    [
        "wrong_field",
        "wrong_row",
        "wrong_header",
        "wrong_sheet",
        "zero_delta",
        "invalid_delta",
        "missing_pair",
    ],
)
def test_invalid_recipes_reject_before_publication(tmp_path, change):
    base = make_base(tmp_path / "original", "xlsx")
    target, distractor = base.manifest.variant_recipes
    if change == "wrong_field":
        target = replace(target, field="Amount")
    elif change == "wrong_row":
        target = replace(target, row=4)
    elif change == "wrong_header":
        target = replace(target, header_row=4, rows=(5, 6), row=5)
    elif change == "wrong_sheet":
        target = replace(target, sheet="absent")
    elif change == "zero_delta":
        target = replace(target, numeric_delta="0")
    elif change == "invalid_delta":
        target = replace(target, numeric_delta="1_000")
    recipes = (target,) if change == "missing_pair" else (target, distractor)
    base = replace(base, manifest=replace(base.manifest, variant_recipes=recipes))
    with pytest.raises(CaseError):
        generate((base,), 7, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_two_bespoke_bases_and_long_ids(tmp_path):
    csv_base = make_base(tmp_path / "csv", "csv")
    xlsx_base = make_base(tmp_path / "xlsx", "xlsx")
    csv_base = replace(csv_base, manifest=replace(csv_base.manifest, case_id="a" * 64))
    suite = generate((xlsx_base, csv_base), 9, tmp_path / "out")
    assert len(suite.cases) == 6
    assert len(suite.families) == 2
    assert all(len(ref.case_id) <= 64 for ref in suite.cases)
    assert len({ref.case_id for ref in suite.cases}) == 6
    for case in frozen_by_id(suite, tmp_path / "out").values():
        assert case.validated.spec.metric_id == "custom_total"
        assert case.gold.origin == "generated_contract"
        if case.validated.manifest.inputs[0].format == "xlsx":
            if case.validated.inputs.snapshots["work"] != xlsx_base.inputs.snapshots["work"]:
                with zipfile.ZipFile(
                    io.BytesIO(case.validated.inputs.snapshots["work"])
                ) as archive:
                    assert all(
                        info.date_time == (1980, 1, 1, 0, 0, 0) for info in archive.infolist()
                    )


def test_empty_duplicate_bases_and_existing_destination(tmp_path):
    base = make_base(tmp_path / "original")
    for bases in ((), (base, base)):
        with pytest.raises(CaseError):
            generate(bases, 7, tmp_path / "bad")
        assert not (tmp_path / "bad").exists()
    existing = tmp_path / "existing"
    existing.mkdir()
    (existing / "keep").write_text("unchanged")
    with pytest.raises(CaseError):
        generate((base,), 7, existing)
    assert (existing / "keep").read_text() == "unchanged"


def test_numeric_xlsx_text_retains_exact_value_and_kind(tmp_path):
    base = make_base(tmp_path / "original", "xlsx", amount="1.000000000000499999", places=12)
    target, distractor = base.manifest.variant_recipes
    base = replace(
        base,
        manifest=replace(
            base.manifest, variant_recipes=(replace(target, numeric_delta="1"), distractor)
        ),
    )
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    target_case = cases[suite.families[0].target]
    assert target_case.gold.expected.value == "2.000000000000"
    changed = target_case.validated.inputs.tables["chosen"].rows[0][1]
    assert (changed.kind, changed.value) == ("text", "2.000000000000499999")


def test_csv_mutation_retains_unrelated_raw_bytes(tmp_path):
    root = tmp_path / "original"
    base = make_base(root)
    data = b'\xef\xbb\xbfflag,amount,note\r\nyes,"100","target"\r\nno,863,"multi\r\nline"\r\n'
    (root / "input.csv").write_bytes(data)
    manifest = replace(
        base.manifest, inputs=(replace(base.manifest.inputs[0], sha256=sha256_bytes(data)),)
    )
    base = validate_case(base.spec, manifest, root)
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    assert cases[suite.families[0].target].validated.inputs.snapshots["work"] == (
        b'\xef\xbb\xbfflag,amount,note\r\nyes,137,"target"\r\nno,863,"multi\r\nline"\r\n'
    )
    assert cases[suite.families[0].distractor].validated.inputs.snapshots["work"] == (
        b'\xef\xbb\xbfflag,amount,note\r\nyes,"100","target"\r\nno,900,"multi\r\nline"\r\n'
    )


def test_each_recipe_is_observable_and_families_share_single_variants(tmp_path):
    base = make_base(tmp_path / "original", "xlsx")
    target, distractor = base.manifest.variant_recipes
    recipes = (
        target,
        distractor,
        replace(target, numeric_delta="11"),
        replace(distractor, numeric_delta="22"),
    )
    base = replace(base, manifest=replace(base.manifest, variant_recipes=recipes))
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    assert len(cases) == 5
    assert len(suite.families) == 4
    assert len({family.base for family in suite.families}) == 1
    assert {cases[family.target].gold.expected.value for family in suite.families} == {"137", "111"}
    assert all(cases[family.distractor].gold.expected.value == "100" for family in suite.families)
    invisible = replace(target, numeric_delta="0.00001")
    base = replace(base, manifest=replace(base.manifest, variant_recipes=recipes + (invisible,)))
    with pytest.raises(CaseError, match="INVISIBLE_TARGET_MUTATION"):
        generate((base,), 7, tmp_path / "must-not-exist")
    assert not (tmp_path / "must-not-exist").exists()


def test_no_recipes_and_header_five_distractor_shuffle(tmp_path):
    base = make_base(tmp_path / "original", "xlsx")
    base_only = replace(base, manifest=replace(base.manifest, variant_recipes=()))
    suite = generate((base_only,), 7, tmp_path / "base-only")
    assert len(suite.cases) == 1 and not suite.families
    shuffle = recipe("column_shuffle", "Distractor", 5, (6, 6), None, None, None)
    base = replace(base, manifest=replace(base.manifest, variant_recipes=(shuffle,)))
    suite = generate((base,), 7, tmp_path / "shuffled")
    cases = frozen_by_id(suite, tmp_path / "shuffled")
    assert len(cases) == 2 and not suite.families
    (changed,) = [
        case
        for case in cases.values()
        if case.validated.inputs.snapshots["work"] != base.inputs.snapshots["work"]
    ]
    book = load_workbook(io.BytesIO(changed.validated.inputs.snapshots["work"]))
    headers = [book["Distractor"].cell(5, col).value for col in range(1, 4)]
    assert headers != ["note", "amount", "flag"]
    assert {
        header: book["Distractor"].cell(6, index + 1).value for index, header in enumerate(headers)
    } == {"note": "decoy", "amount": "863", "flag": "no"}
    assert book["Distractor"]["A1"].value == "Keep this preamble"
    assert book["Chosen"]["B2"].value == "100"
    book.close()


def test_recipe_budget_and_seed_are_checked_before_publication(tmp_path):
    base = make_base(tmp_path / "original")
    target, distractor = base.manifest.variant_recipes
    over_budget = replace(
        base, manifest=replace(base.manifest, variant_recipes=(target,) * 99 + (distractor,))
    )
    with pytest.raises(CaseError, match="SUITE_LIMIT"):
        generate((over_budget,), 7, tmp_path / "over-budget")
    for invalid in (True, 1.5, "7"):
        with pytest.raises(CaseError, match="INVALID_SEED"):
            generate((base,), invalid, tmp_path / "invalid-seed")
    assert not (tmp_path / "over-budget").exists()
    assert not (tmp_path / "invalid-seed").exists()


def test_xlsx_numeric_roundtrip_cannot_fabricate_visible_delta(tmp_path):
    root = tmp_path / "original"
    base = make_base(root, "xlsx", places=12)
    book = load_workbook(root / "input.xlsx")
    book["Chosen"]["B2"] = 100
    book.save(root / "input.xlsx")
    book.close()
    data = (root / "input.xlsx").read_bytes()
    target, distractor = base.manifest.variant_recipes
    manifest = replace(
        base.manifest,
        inputs=(replace(base.manifest.inputs[0], sha256=sha256_bytes(data)),),
        variant_recipes=(replace(target, numeric_delta="0.000000000000499999"), distractor),
    )
    base = validate_case(base.spec, manifest, root)
    with pytest.raises(CaseError):
        generate((base,), 7, tmp_path / "out")
    assert not (tmp_path / "out").exists()


def test_legal_implicit_xlsx_coordinates_are_supported(tmp_path):
    root = tmp_path / "original"
    base = make_base(root, "xlsx")
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base.inputs.snapshots["work"])) as source:
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as output:
            for info in source.infolist():
                data = source.read(info)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = re.sub(rb'(<row|<c) r="[A-Z0-9]+"', rb"\1", data)
                output.writestr(info, data)
    data = stream.getvalue()
    (root / "input.xlsx").write_bytes(data)
    manifest = replace(
        base.manifest, inputs=(replace(base.manifest.inputs[0], sha256=sha256_bytes(data)),)
    )
    base = validate_case(base.spec, manifest, root)
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    (family,) = suite.families
    assert [
        cases[key].gold.expected.value for key in (family.base, family.target, family.distractor)
    ] == ["100", "137", "100"]


@pytest.mark.parametrize(
    "operation,want", [("count", ["1", "2", "1"]), ("ratio", ["0.1159", "0.1587", "0.1159"])]
)
def test_count_and_ratio_consume_same_generation_contract(tmp_path, operation, want):
    root = tmp_path / "original"
    base = make_base(root, "xlsx")
    if operation == "count":
        metric = CountSpec("count", "chosen", FilterSpec("flag", "yes"))
        sources = base.spec.sources
        target = recipe("target_filter", "Chosen", row=3, field="flag", text="yes")
        task = {"description": "Count rows 2..3 in Chosen with exact flag yes; unit count."}
    else:
        metric = RatioSpec(
            "ratio",
            AggregateSumSpec("sum", "chosen", "amount", "CNY"),
            AggregateSumSpec("sum", "denominator", "amount", "CNY"),
            4,
            "HALF_UP",
        )
        sources = {
            "chosen": SourceSpec("work", "Chosen", 1, (2, 2)),
            "denominator": SourceSpec("work", "Chosen", 1, (3, 3)),
        }
        target = recipe("target_numeric", "Chosen", rows=(2, 2))
        task = {
            "description": "Divide Chosen row2 amount by row3 amount, CNY/CNY, HALF_UP four places."
        }
    spec = replace(base.spec, metric=metric, sources=sources)
    confirmation = {
        "reviewed_contract_sha256": sha256_bytes(canonical_json(spec)),
        "task_sha256": sha256_bytes(canonical_json(task)),
        "reviewer": "custom-test",
    }
    (root / "task.json").write_bytes(canonical_json(task))
    manifest = replace(
        base.manifest,
        business_confirmation=confirmation,
        variant_recipes=(target, base.manifest.variant_recipes[1]),
    )
    base = validate_case(spec, manifest, root)
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    (family,) = suite.families
    assert [
        cases[key].gold.expected.value for key in (family.base, family.target, family.distractor)
    ] == want


def test_xlsx_opaque_namespace_attributes_survive_normalization(tmp_path):
    root = tmp_path / "original"
    base = make_base(root, "xlsx")
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base.inputs.snapshots["work"])) as source:
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as output:
            for info in source.infolist():
                raw = source.read(info)
                if info.filename == "xl/worksheets/sheet1.xml":
                    raw = raw.replace(
                        b"<worksheet ",
                        b'<worksheet xmlns:custom="urn:custom" '
                        b'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
                        b'mc:Ignorable="custom" ',
                        1,
                    )
                output.writestr(info, raw)
    data = stream.getvalue()
    (root / "input.xlsx").write_bytes(data)
    manifest = replace(
        base.manifest, inputs=(replace(base.manifest.inputs[0], sha256=sha256_bytes(data)),)
    )
    base = validate_case(base.spec, manifest, root)
    suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    with zipfile.ZipFile(
        io.BytesIO(cases[suite.families[0].target].validated.inputs.snapshots["work"])
    ) as archive:
        for part in ("xl/worksheets/sheet1.xml", "docProps/core.xml"):
            raw = archive.read(part)
            prefixes = {
                prefix for _, (prefix, _) in ET.iterparse(io.BytesIO(raw), events=("start-ns",))
            }
            tree = ET.fromstring(raw)
            for element in tree.iter():
                for attribute, value in element.attrib.items():
                    if attribute.endswith("}Ignorable"):
                        assert set(value.split()) <= prefixes
                    elif attribute.endswith("}type") and ":" in value:
                        assert value.split(":", 1)[0] in prefixes


def test_decimal_ambient_context_cannot_change_mutations(tmp_path):
    base = make_base(tmp_path / "original", "xlsx")
    with localcontext() as context:
        context.prec = 2
        context.Emax = 1
        context.Emin = -1
        context.clamp = 1
        context.traps[Inexact] = True
        suite = generate((base,), 7, tmp_path / "out")
    cases = frozen_by_id(suite, tmp_path / "out")
    (family,) = suite.families
    assert [
        cases[key].gold.expected.value for key in (family.base, family.target, family.distractor)
    ] == ["100", "137", "100"]
