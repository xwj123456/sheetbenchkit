"""Deterministic, declared single-location mutations of verified input snapshots.

Gold comes from the reference contract, never independent business preset review.
Every numeric/filter recipe is checked separately before the atomic suite freeze.
"""

from __future__ import annotations

import csv
import io
import random
import re
import tempfile
import zipfile
from collections.abc import Mapping
from dataclasses import replace
from decimal import Decimal, localcontext
from pathlib import Path
from typing import Any

from defusedxml.minidom import parseString  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.cell import (  # type: ignore[import-untyped]
    column_index_from_string,
    get_column_letter,
)

from . import models as m
from .artifacts import canonical_json, freeze_suite, sha256_bytes
from .contracts import to_document, validate_case, validate_documents
from .loaders import MAX_FILE_BYTES, _csv_tables, _xlsx_tables
from .reference import evaluate_reference

_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_CORE = "{http://purl.org/dc/terms/}"
_NUMERIC = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE]([+-]?[0-9]+))?\Z")


def _decimal(raw: Any) -> Decimal:
    if not isinstance(raw, str):
        raise m.CaseError("INVALID_RECIPE", "numeric text required")
    match = _NUMERIC.fullmatch(raw)
    if not match or len(raw) > 80:
        raise m.CaseError("INVALID_RECIPE", "invalid numeric literal")
    mantissa = re.split("[eE]", raw)[0].lstrip("+-")
    if (
        len(mantissa.replace(".", "").lstrip("0") or "0") > 30
        or len(mantissa.split(".", 1)[1] if "." in mantissa else "") > 30
        or not -12 <= int(match.group(1) or "0") <= 12
    ):
        raise m.CaseError("INVALID_RECIPE", "numeric literal exceeds approved bounds")
    return Decimal(raw)


def _numeric_change(cell: m.Cell, delta: str | None) -> str:
    if cell.kind not in {"text", "number"}:
        raise m.CaseError("INVALID_RECIPE", "numeric mutation requires a finite numeric cell")
    original, change = _decimal(cell.value), _decimal(delta)
    if not change:
        raise m.CaseError("INVALID_RECIPE", "zero numeric change")
    with localcontext() as context:
        context.prec = 200  # Input exponent/coefficient bounds make this sum exact.
        context.Emax, context.Emin, context.clamp = 100, -100, 0
        result = original + change
    parts = result.as_tuple()
    exponent = parts.exponent
    assert isinstance(exponent, int)
    digits = "".join(str(digit) for digit in parts.digits)
    if not result:
        return "0"
    while digits.endswith("0"):
        digits = digits[:-1]
        exponent += 1
    # Retain the legal scientific-exponent range rather than expanding tiny
    # legal inputs into a forbidden >30-place raw fractional representation.
    explicit = min(12, max(-12, exponent))
    remaining = exponent - explicit
    if remaining >= 0:
        mantissa = digits + "0" * remaining
    else:
        split = len(digits) + remaining
        mantissa = (
            (digits[:split] + "." + digits[split:]) if split > 0 else ("0." + "0" * -split + digits)
        )
    raw = ("-" if parts.sign else "") + mantissa + (f"e{explicit}" if explicit else "")
    _decimal(raw)
    # Use ordinary decimal text for common values, preserving bounds for extremes.
    plain = format(result, "f")
    try:
        _decimal(plain)
        return plain
    except m.CaseError:
        return raw


def _recipe_table(data: bytes, format_: str, recipe: m.MutationRecipe) -> m.Table:
    if recipe.rows[1] - recipe.rows[0] + 1 > 1000:
        raise m.CaseError("INVALID_RECIPE", "row budget")
    source = m.SourceSpec(recipe.file_id, recipe.sheet, recipe.header_row, recipe.rows)
    if format_ == "csv" and recipe.sheet != "@csv":
        raise m.CaseError("INVALID_RECIPE", "CSV sheet must be @csv")
    tables, limits, _reads = (
        _csv_tables(data, {"recipe": source})
        if format_ == "csv"
        else _xlsx_tables(data, {"recipe": source})
    )
    if limits or "recipe" not in tables:
        raise m.CaseError("INVALID_RECIPE", {"limits": limits})
    table = tables["recipe"]
    if (
        not table.headers
        or any(not isinstance(field, str) or not field for field in table.headers)
        or len(set(table.headers)) != len(table.headers)
    ):
        raise m.CaseError("INVALID_RECIPE", "unique nonempty exact headers required")
    if recipe.kind != "column_shuffle" and table.headers.count(recipe.field) != 1:
        raise m.CaseError("INVALID_RECIPE", "exact field missing or ambiguous")
    return table


def _target_location(base: m.ValidatedCase, recipe: m.MutationRecipe) -> None:
    if not recipe.kind.startswith("target_"):
        return
    metric = base.spec.metric
    aggregates = (
        (metric.numerator, metric.denominator) if isinstance(metric, m.RatioSpec) else (metric,)
    )
    for aggregate in aggregates:
        source = base.spec.sources[aggregate.source]
        if (source.file_id, source.sheet, source.header_row, source.rows) != (
            recipe.file_id,
            recipe.sheet,
            recipe.header_row,
            recipe.rows,
        ):
            continue
        if recipe.kind == "target_numeric" and isinstance(aggregate, m.AggregateSumSpec):
            if recipe.field == aggregate.field:
                return
        if recipe.kind == "target_filter" and aggregate.filter is not None:
            if recipe.field == aggregate.filter.field:
                return
    raise m.CaseError("INVALID_RECIPE", "target is not a declared metric/filter location")


def _permutation(width: int, seed: int, recipe: m.MutationRecipe) -> list[int]:
    if width < 2:
        raise m.CaseError("INVALID_RECIPE", "shuffle requires at least two columns")
    order = list(range(width))
    rng = random.Random(sha256_bytes(canonical_json({"seed": seed, "recipe": recipe})))
    rng.shuffle(order)
    if order == list(range(width)):
        order = order[1:] + order[:1]
    return order


def _csv_mutation(data: bytes, recipe: m.MutationRecipe, table: m.Table, seed: int) -> bytes:
    text = data.decode("utf-8-sig")
    # Grammar was validated by _recipe_table. Record exact field spans, including
    # quoted multiline text, so unrelated BOM/newline/quoting bytes remain intact.
    records: list[list[tuple[int, int]]] = []
    record: list[tuple[int, int]] = []
    start = position = 0
    while position < len(text):
        if text[position] == '"':
            position += 1
            while True:
                quote = text.index('"', position)
                position = quote + 1
                if position < len(text) and text[position] == '"':
                    position += 1
                else:
                    break
        else:
            while position < len(text) and text[position] not in ",\r\n":
                position += 1
        record.append((start, position))
        if position == len(text):
            records.append(record)
            break
        delimiter = text[position]
        position += 1
        if delimiter != ",":
            if delimiter == "\r" and position < len(text) and text[position] == "\n":
                position += 1
            records.append(record)
            record = []
        start = position
    if record and position == len(text) and start == len(text):
        # A final comma denotes a trailing blank field, even without a newline.
        record.append((position, position))
        records.append(record)
    replacements: list[tuple[int, int, str]] = []
    if recipe.kind == "column_shuffle":
        order = _permutation(len(table.headers), seed, recipe)
        for row in (recipe.header_row, *range(recipe.rows[0], recipe.rows[1] + 1)):
            spans = records[row - 1]
            replacements.extend(
                (left, right, text[spans[old][0] : spans[old][1]])
                for (left, right), old in zip(spans, order, strict=True)
            )
    else:
        assert recipe.row is not None and recipe.field is not None
        column = table.headers.index(recipe.field)
        cell = table.rows[recipe.row - recipe.rows[0]][column]
        replacement = _replacement(cell, recipe)
        stream = io.StringIO(newline="")
        csv.writer(stream, lineterminator="\r\n").writerow((replacement,))
        left, right = records[recipe.row - 1][column]
        replacements.append((left, right, stream.getvalue()[:-2]))
    for left, right, replacement in sorted(replacements, reverse=True):
        text = text[:left] + replacement + text[right:]
    return (b"\xef\xbb\xbf" if data.startswith(b"\xef\xbb\xbf") else b"") + text.encode("utf-8")


def _replacement(cell: m.Cell, recipe: m.MutationRecipe) -> str:
    if recipe.kind.endswith("_numeric"):
        return _numeric_change(cell, recipe.numeric_delta)
    if cell.kind not in {"text", "blank"} or recipe.text_replacement is None:
        raise m.CaseError("INVALID_RECIPE", "filter mutation requires a text/blank cell")
    if (cell.value or "") == recipe.text_replacement:
        raise m.CaseError("INVALID_RECIPE", "unchanged filter text")
    return recipe.text_replacement


def _children(node: Any, name: str) -> list[Any]:
    return [
        child
        for child in node.childNodes
        if child.nodeType == child.ELEMENT_NODE
        and child.namespaceURI == _NS[1:-1]
        and child.localName == name
    ]


def _xml_element(document: Any, context: Any, name: str) -> Any:
    prefix = context.prefix + ":" if context.prefix else ""
    return document.createElementNS(_NS[1:-1], prefix + name)


def _xlsx_mutation(data: bytes, recipe: m.MutationRecipe, table: m.Table, seed: int) -> bytes:
    # _recipe_table already checked ZIP expansion/XML safety before this parser.
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        part = book[recipe.sheet]._worksheet_path.lstrip("/")
    finally:
        book.close()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        documents = {info.filename: source.read(info) for info in source.infolist()}
    # DOM retains aliases used inside opaque QName attributes, without global state.
    document = parseString(documents[part], forbid_dtd=True)
    sheet_nodes = _children(document.documentElement, "sheetData")
    if len(sheet_nodes) != 1:
        raise m.CaseError("INVALID_RECIPE", "no sheet data")
    sheet_data = sheet_nodes[0]
    rows = {}
    previous_row = 0
    for physical_row in _children(sheet_data, "row"):
        number = int(physical_row.getAttribute("r") or previous_row + 1)
        physical_row.setAttribute("r", str(number))
        rows[number] = physical_row
        previous_row = number
        previous_column = 0
        for physical_cell in _children(physical_row, "c"):
            address = physical_cell.getAttribute("r")
            column = (
                column_index_from_string(address.rstrip("0123456789"))
                if address
                else previous_column + 1
            )
            physical_cell.setAttribute("r", f"{get_column_letter(column)}{number}")
            previous_column = column
    if recipe.kind == "column_shuffle":
        order = _permutation(len(table.headers), seed, recipe)
        for number in (recipe.header_row, *range(recipe.rows[0], recipe.rows[1] + 1)):
            shuffle_row = rows.get(number)
            if shuffle_row is None:
                continue  # Sparse absent blank rows remain absent blank rows.
            cells = {cell.getAttribute("r"): cell for cell in _children(shuffle_row, "c")}
            for cell in cells.values():
                shuffle_row.removeChild(cell)
            for new_column, old_column in enumerate(order, 1):
                permuted_cell = cells.get(f"{get_column_letter(old_column + 1)}{number}")
                if permuted_cell is not None:
                    permuted_cell.setAttribute("r", f"{get_column_letter(new_column)}{number}")
                    shuffle_row.appendChild(permuted_cell)
    else:
        assert recipe.row is not None and recipe.field is not None
        column = table.headers.index(recipe.field)
        cell_value = table.rows[recipe.row - recipe.rows[0]][column]
        replacement = _replacement(cell_value, recipe)
        address = f"{get_column_letter(column + 1)}{recipe.row}"
        mutation_row = rows.get(recipe.row)
        if mutation_row is None:
            mutation_row = _xml_element(document, sheet_data, "row")
            mutation_row.setAttribute("r", str(recipe.row))
            sheet_data.appendChild(mutation_row)
            for row in sorted(
                _children(sheet_data, "row"), key=lambda node: int(node.getAttribute("r"))
            ):
                sheet_data.appendChild(row)
        mutation_cell = next(
            (cell for cell in _children(mutation_row, "c") if cell.getAttribute("r") == address),
            None,
        )
        if mutation_cell is None:
            mutation_cell = _xml_element(document, mutation_row, "c")
            mutation_cell.setAttribute("r", address)
            mutation_row.appendChild(mutation_cell)
        for child in list(mutation_cell.childNodes):
            mutation_cell.removeChild(child)
        if recipe.kind.endswith("_numeric") and cell_value.kind == "number":
            mutation_cell.setAttribute("t", "n")
            value = _xml_element(document, mutation_cell, "v")
            value.appendChild(document.createTextNode(replacement))
            mutation_cell.appendChild(value)
        else:
            mutation_cell.setAttribute("t", "inlineStr")
            inline = _xml_element(document, mutation_cell, "is")
            text = _xml_element(document, mutation_cell, "t")
            text.setAttributeNS("http://www.w3.org/XML/1998/namespace", "xml:space", "preserve")
            text.appendChild(document.createTextNode(replacement))
            inline.appendChild(text)
            mutation_cell.appendChild(inline)
        for cell in sorted(
            _children(mutation_row, "c"),
            key=lambda node: (
                len(node.getAttribute("r").rstrip("0123456789")),
                node.getAttribute("r"),
            ),
        ):
            mutation_row.appendChild(cell)
    documents[part] = document.toxml(encoding="utf-8")
    document.unlink()
    if "docProps/core.xml" in documents:
        core = parseString(documents["docProps/core.xml"], forbid_dtd=True)
        for key in ("created", "modified"):
            for element in core.getElementsByTagNameNS(_CORE[1:-1], key):
                for child in list(element.childNodes):
                    element.removeChild(child)
                element.appendChild(core.createTextNode("2000-01-01T00:00:00Z"))
        documents["docProps/core.xml"] = core.toxml(encoding="utf-8")
        core.unlink()
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name in sorted(documents):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, documents[name])
    return stream.getvalue()


def _identity(base: m.ValidatedCase, suffix: str) -> str:
    digest = sha256_bytes(canonical_json({"base": base.manifest.case_id, "suffix": suffix}))[:16]
    return base.manifest.case_id[:34] + "_" + suffix[:12] + "_" + digest


def _revalidate(
    base: m.ValidatedCase, snapshots: Mapping[str, bytes], case_id: str
) -> m.ValidatedCase:
    entries = tuple(
        replace(
            entry,
            relative_path=f"inputs/{entry.file_id}.{entry.format}",
            sha256=None if entry.expected_missing else sha256_bytes(snapshots[entry.file_id]),
        )
        for entry in base.manifest.inputs
    )
    manifest = replace(base.manifest, case_id=case_id, inputs=entries)
    # All temporary reads use copies of validated snapshots, outside the source tree.
    with tempfile.TemporaryDirectory(prefix="sheetbenchkit-variants-") as temporary:
        root = Path(temporary)
        (root / "inputs").mkdir()
        for entry in entries:
            if not entry.expected_missing:
                (root / entry.relative_path).write_bytes(snapshots[entry.file_id])
        (root / "task.json").write_bytes(canonical_json(base.task))
        return validate_case(base.spec, manifest, root)


def generate_suite(
    base_cases: tuple[m.ValidatedCase, ...], seed: int, destination: Path
) -> m.Suite:
    """Revalidate actual single-recipe files, check relations, then freeze once.

    Each base has one case; each recipe has its own variant. All target/distractor
    pairs form families sharing their base and independently checked variants.
    A base with no such recipes (including shuffle-only bases) needs no family.
    """
    if type(seed) is not int:
        raise m.CaseError("INVALID_SEED", "integer required")
    if destination.exists() or destination.is_symlink():
        raise m.CaseError("DESTINATION_EXISTS", str(destination))
    if not base_cases or sum(1 + len(base.manifest.variant_recipes) for base in base_cases) > 100:
        raise m.CaseError("SUITE_LIMIT", "generated suite must contain 1..100 cases")
    ids = [base.manifest.case_id for base in base_cases]
    if len(ids) != len(set(ids)):
        raise m.CaseError("DUPLICATE_CASE_ID")
    candidates: list[m.CandidateCase] = []
    families: list[m.FamilySpec] = []
    for base in sorted(base_cases, key=lambda case: case.manifest.case_id):
        validate_documents(base.spec, base.manifest)
        targets = [r for r in base.manifest.variant_recipes if r.kind.startswith("target_")]
        distractors = [r for r in base.manifest.variant_recipes if r.kind.startswith("distractor_")]
        if bool(targets) != bool(distractors):
            raise m.CaseError("INVALID_RECIPE", "target/distractor family requires both kinds")
        snapshots = dict(base.inputs.snapshots)
        known = {entry.file_id for entry in base.manifest.inputs}
        if set(snapshots) - known:
            raise m.CaseError("UNKNOWN_SNAPSHOT")
        for entry in base.manifest.inputs:
            data = snapshots.get(entry.file_id)
            if entry.expected_missing:
                if data is not None:
                    raise m.CaseError("EXPECTED_MISSING_PRESENT", entry.file_id)
            elif data is None or len(data) > MAX_FILE_BYTES or sha256_bytes(data) != entry.sha256:
                raise m.CaseError("INPUT_HASH_MISMATCH", "full verified <=10MiB snapshot required")
        verified = _revalidate(base, snapshots, _identity(base, "base"))
        expected = evaluate_reference(verified)
        candidates.append(
            m.CandidateCase(
                verified,
                expected,
                "generated_contract",
                base.manifest.business_confirmation,
                {"base_case_id": base.manifest.case_id, "kind": "base", "seed": seed},
            )
        )
        target_ids: list[tuple[str, m.MutationRecipe]] = []
        distractor_ids: list[tuple[str, m.MutationRecipe]] = []
        entries_by_id = {entry.file_id: entry for entry in base.manifest.inputs}
        for index, recipe in enumerate(base.manifest.variant_recipes):
            _target_location(base, recipe)
            data = snapshots.get(recipe.file_id)
            if data is None:
                raise m.CaseError("INVALID_RECIPE", "mutation requires a verified snapshot")
            format_ = entries_by_id[recipe.file_id].format
            table = _recipe_table(data, format_, recipe)
            changed = (
                _csv_mutation(data, recipe, table, seed)
                if format_ == "csv"
                else _xlsx_mutation(data, recipe, table, seed)
            )
            if recipe.kind != "column_shuffle":
                assert recipe.row is not None and recipe.field is not None
                readback = _recipe_table(changed, format_, recipe)
                column = table.headers.index(recipe.field)
                old = table.rows[recipe.row - recipe.rows[0]][column]
                actual = readback.rows[recipe.row - recipe.rows[0]][column]
                intended = _replacement(old, recipe)
                if recipe.kind.endswith("_numeric"):
                    if actual.kind != old.kind or _decimal(actual.value) != _decimal(intended):
                        raise m.CaseError(
                            "INVALID_RECIPE", "numeric mutation cannot roundtrip exactly"
                        )
                elif (actual.value or "") != intended or actual.kind not in {"text", "blank"}:
                    raise m.CaseError("INVALID_RECIPE", "filter text cannot roundtrip exactly")
            if changed == data or len(changed) > MAX_FILE_BYTES:
                raise m.CaseError("INVALID_RECIPE", "unchanged or oversized mutation")
            altered = {**snapshots, recipe.file_id: changed}
            variant = _revalidate(base, altered, _identity(base, f"variant_{index}"))
            gold = evaluate_reference(variant)
            if recipe.kind.startswith("target_"):
                if (
                    expected.status != "VALUE"
                    or gold.status != "VALUE"
                    or gold.value == expected.value
                ):
                    raise m.CaseError("INVISIBLE_TARGET_MUTATION", to_document(recipe))
                target_ids.append((variant.manifest.case_id, recipe))
            else:
                if gold != expected:
                    raise m.CaseError("DISTRACTOR_CHANGED_GOLD", to_document(recipe))
                if recipe.kind.startswith("distractor_"):
                    distractor_ids.append((variant.manifest.case_id, recipe))
            candidates.append(
                m.CandidateCase(
                    variant,
                    gold,
                    "generated_contract",
                    base.manifest.business_confirmation,
                    {
                        "base_case_id": base.manifest.case_id,
                        "kind": recipe.kind,
                        "recipe": to_document(recipe),
                        "seed": seed,
                    },
                )
            )
        for target_index, (target_id, target) in enumerate(target_ids):
            for distractor_index, (distractor_id, distractor) in enumerate(distractor_ids):
                families.append(
                    m.FamilySpec(
                        _identity(base, f"f_{target_index}_{distractor_index}"),
                        verified.manifest.case_id,
                        target_id,
                        distractor_id,
                        {"target": to_document(target), "distractor": to_document(distractor)},
                    )
                )
    return freeze_suite(tuple(candidates), tuple(families), destination)
