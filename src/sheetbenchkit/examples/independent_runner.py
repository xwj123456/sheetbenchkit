"""A small independent offline adapter that reads the envelope's absolute files.

This example uses openpyxl/stdlib directly, and does not import the toolkit's
loaders, reference, variants, grader or arithmetic. Its input is a preflighted
TaskEnvelope. It is trusted local code, not a sandbox or a natural-language agent.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

_NUMBER = re.compile(r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE]([+-]?[0-9]+))?\Z")


@dataclass(frozen=True)
class ReadCell:
    """Retain intrinsic input kind separately from literal spelling."""

    kind: str
    value: Any


def xlsx_cell(cell: Any) -> ReadCell:
    value = cell.value
    if value is None or value == "":
        return ReadCell("blank", None)
    if cell.data_type == "f":
        return ReadCell("formula", value)
    if cell.data_type == "e":
        return ReadCell("error", value)
    if isinstance(value, bool):
        return ReadCell("bool", value)
    if isinstance(value, (date, datetime, time, timedelta)):
        return ReadCell("date", value)
    if isinstance(value, (int, float, Decimal)):
        return ReadCell("number", str(Decimal(str(value))))
    if isinstance(value, str):
        return ReadCell("text", value)
    raise ValueError("Unknown XLSX cell representation")


class Refusal(Exception):
    """A legitimate business boundary in an otherwise valid task."""


def declared_bindings(contract: dict[str, Any]) -> list[dict[str, Any]]:
    metric = contract["metric"]
    items = (metric["numerator"], metric["denominator"]) if metric["op"] == "ratio" else (metric,)
    needed: dict[str, set[str]] = {}
    for item in items:
        fields = needed.setdefault(item["source"], set())
        if item["op"] == "sum":
            fields.add(item["field"])
        if item.get("filter") is not None:
            fields.add(item["filter"]["field"])
    return [
        {"source": source, **contract["sources"][source], "fields": sorted(fields)}
        for source, fields in sorted(needed.items())
    ]


def rounded(value: Fraction, places: int) -> str:
    """Exact integer long-division rounding, without a Decimal context."""
    scaled, remainder = divmod(abs(value.numerator) * 10**places, value.denominator)
    scaled += int(remainder * 2 >= value.denominator)
    sign = "-" if value < 0 and scaled else ""
    digits = str(scaled).zfill(places + 1)
    return sign + (digits[:-places] + "." + digits[-places:] if places else digits)


class Calculator:
    def __init__(self, envelope: dict[str, Any]):
        self.envelope = envelope
        self.contract = envelope["output_protocol"]["contract"]
        self.labels = envelope["output_protocol"]["unit_labels"]
        self.entries = {item["file_id"]: item for item in envelope["inputs"]}
        self.tables: dict[str, tuple[list[Any], list[list[ReadCell]]]] = {}

    def table(self, source_id: str) -> tuple[list[Any], list[list[ReadCell]]]:
        if source_id in self.tables:
            return self.tables[source_id]
        source = self.contract["sources"][source_id]
        entry = self.entries[source["file_id"]]
        path = Path(entry["relative_path"])
        if not path.is_absolute():
            raise ValueError("TaskEnvelope input paths must be absolute")
        if entry["expected_missing"]:
            raise Refusal("MISSING_SOURCE")
        start, stop = source["rows"]
        header = source["header_row"]
        if stop - start + 1 > 1000:
            raise Refusal("UNSUPPORTED_INPUT_LIMIT")
        if entry["format"] == "csv":
            if source["sheet"] != "@csv":
                raise Refusal("MISSING_SOURCE")
            with path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, strict=True))
            headers = rows[header - 1]
            if len(headers) > 100:
                raise Refusal("UNSUPPORTED_INPUT_LIMIT")
            data = [
                [ReadCell("text", value) if value else ReadCell("blank", None) for value in row]
                for row in rows[start - 1 : stop]
            ]
        else:
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                if source["sheet"] not in book.sheetnames:
                    raise Refusal("MISSING_SOURCE")
                sheet = book[source["sheet"]]
                sheet.reset_dimensions()
                # Read header and declared data separately; no physical row after
                # the workbook's end is required for a declared blank row to exist.
                header_rows = list(sheet.iter_rows(min_row=header, max_row=header))
                headers = [cell.value for cell in header_rows[0]] if header_rows else []
                if len(headers) > 100:
                    raise Refusal("UNSUPPORTED_INPUT_LIMIT")
                if not headers:
                    raise Refusal("MISSING_FIELD")
                data = []
                for row in sheet.iter_rows(min_row=start, max_row=stop):
                    if len(row) > 100:
                        raise Refusal("UNSUPPORTED_INPUT_LIMIT")
                    values = [xlsx_cell(cell) for cell in row]
                    values.extend(
                        ReadCell("blank", None) for _ in range(len(headers) - len(values))
                    )
                    data.append(values)
                data.extend(
                    [ReadCell("blank", None) for _ in headers]
                    for _ in range(stop - start + 1 - len(data))
                )
            finally:
                book.close()
        result = headers, data
        self.tables[source_id] = result
        return result

    def field(self, headers: list[Any], name: str) -> int:
        matches = [index for index, item in enumerate(headers) if item == name]
        if not matches:
            raise Refusal("MISSING_FIELD")
        if len(matches) > 1:
            raise Refusal("AMBIGUOUS_FIELD")
        return matches[0]

    def number(self, value: Any) -> Fraction:
        if value is None or not str(value).strip():
            raise Refusal("MISSING_VALUE")
        raw = str(value)
        match = _NUMBER.fullmatch(raw)
        if not match:
            raise Refusal("INVALID_VALUE")
        if len(raw) > 80:
            raise Refusal("UNSUPPORTED_INPUT_LIMIT")
        mantissa = re.split("[eE]", raw)[0].lstrip("+-")
        if (
            len(mantissa.replace(".", "").lstrip("0") or "0") > 30
            or len(mantissa.split(".", 1)[1] if "." in mantissa else "") > 30
            or not -12 <= int(match.group(1) or "0") <= 12
        ):
            raise Refusal("UNSUPPORTED_INPUT_LIMIT")
        return Fraction(Decimal(raw))

    def values(self, item: dict[str, Any]) -> list[Fraction]:
        headers, rows = self.table(item["source"])
        column = self.field(headers, item["field"]) if item["op"] == "sum" else None
        condition = item.get("filter")
        if condition is not None:
            chosen = self.field(headers, condition["field"])
            selected = []
            for row in rows:
                cell = row[chosen]
                if cell.kind == "blank":
                    continue
                if cell.kind != "text" or not isinstance(cell.value, str):
                    raise Refusal("INVALID_FILTER_VALUE")
                if not cell.value.strip():
                    continue
                if cell.value == condition["eq"]:
                    selected.append(row)
            rows = selected
        if item["op"] == "count":
            return [Fraction(1) for _ in rows]
        source = self.contract["sources"][item["source"]]
        label = self.labels.get(source["file_id"], {}).get(source["sheet"], {}).get(item["field"])
        if label != item["unit"]:
            raise Refusal("UNSUPPORTED_UNIT_CONVERSION")
        assert column is not None
        values = []
        for row in rows:
            cell = row[column]
            if cell.kind not in {"blank", "number", "text"}:
                raise Refusal("UNSUPPORTED_CELL")
            values.append(self.number(cell.value))
        return values

    def ratio(self, numerator: list[Fraction], denominator: list[Fraction]) -> Fraction:
        divisor = sum(denominator, Fraction(0))
        if not divisor:
            raise Refusal("ZERO_DENOMINATOR")
        return sum(numerator, Fraction(0)) / divisor

    def compatible(self, numerator: dict[str, Any], denominator: dict[str, Any]) -> bool:
        return numerator.get("unit", "count") == denominator.get("unit", "count")

    def format(self, value: Fraction, places: int) -> str:
        return rounded(value, places)

    def calculate(self) -> tuple[str, str]:
        metric = self.contract["metric"]
        if metric["op"] == "ratio":
            numerator, denominator = metric["numerator"], metric["denominator"]
            top, bottom = self.values(numerator), self.values(denominator)
            if not self.compatible(numerator, denominator):
                raise Refusal("UNSUPPORTED_UNIT_CONVERSION")
            return self.format(self.ratio(top, bottom), metric["places"]), "ratio"
        total = sum(self.values(metric), Fraction(0))
        if metric["op"] == "count":
            return str(total.numerator), "count"
        return self.format(total, metric["places"]), metric["unit"]

    def solve(self) -> dict[str, Any]:
        result = {
            "schema_version": "1",
            "case_id": self.envelope["case_id"],
            "metric_id": self.envelope["metric_id"],
            "bindings": declared_bindings(self.contract),
        }
        try:
            value, unit = self.calculate()
            result.update(status="VALUE", value=value, unit=unit, reason=None)
        except Refusal as refusal:
            result.update(status="ABSTAIN", value=None, unit=None, reason=str(refusal))
        return result


def solve(envelope: dict[str, Any]) -> dict[str, Any]:
    return Calculator(envelope).solve()


def main() -> int:
    result = solve(json.load(sys.stdin))
    print(json.dumps(result, ensure_ascii=False, allow_nan=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
